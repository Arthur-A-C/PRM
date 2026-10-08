"""Template-driven CIGAM workbook export; no ERP write path."""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

from app.schemas.registration import RegistrationProduct


KNOWN_FIELDS = {
    "Modelo": lambda product: product.model,
    "Aplicação Material - Descrição": lambda product: product.technical_name,
    "Codigo Cigam": lambda product: product.cigam_code,
}


def export_cigam_xlsx(product: RegistrationProduct, template: bytes) -> tuple[bytes, list[str]]:
    source = load_workbook(io.BytesIO(template), read_only=True, data_only=False)
    try:
        if len(source.worksheets) != 1:
            raise ValueError("Modelo CIGAM deve ter uma única aba validada.")
        headers = [str(value).strip() if value is not None else "" for value in next(source.active.iter_rows(max_row=1, values_only=True))]
    finally:
        source.close()
    if len(headers) != len(set(headers)) or not {"Codigo Cigam", "Modelo"}.issubset(headers):
        raise ValueError("Cabeçalhos CIGAM incompatíveis ou duplicados.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cadastro"
    sheet.append(headers)
    warnings = []
    row = []
    for header in headers:
        getter = KNOWN_FIELDS.get(header)
        value = getter(product) if getter else product.cigam_fields.get(header)
        if value is None or str(value).strip() == "":
            warnings.append(f"Campo CIGAM pendente: {header}")
            row.append(None)
        else:
            text = str(value).strip()
            if text.startswith(("=", "+", "-", "@")):
                raise ValueError(f"Valor CIGAM inseguro em {header}.")
            row.append(text)
    sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue(), warnings
