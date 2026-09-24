from __future__ import annotations

import hashlib
import html
import json
import math
import re
import unicodedata
import uuid
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import Settings
from app.models import (
    Brand,
    BrandAlias,
    CatalogImport,
    Category,
    Company,
    ImportRow,
    Offer,
    OfferEnrichment,
    OfferObservation,
    OfferPrice,
)
from app.services.catalog_treatment import product_content_sha256, treat_raw


class CatalogValidationError(ValueError):
    pass


FIELD_ALIASES = {
    "external_code": ("codigo", "cod material", "c d material"),
    "name": ("nome", "descricao material", "descri o material"),
    "source_url": ("link", "url"),
    "brand": ("marca", "fabricante"),
    "description": ("descricao", "descri o"),
    "application": ("aplicacao material", "aplica o material"),
    "pix_price": ("preco pix",),
    "card_price": ("preco cartao",),
    "sale_price": ("preco venda", "pre o venda"),
    "availability": ("estoque",),
    "stock": ("estoque atual",),
    "category": ("desc grupo",),
    "subcategory": ("desc subgrupo",),
    "manufacturer_code": ("cod fabrica", "c d f brica", "c d fabrica"),
    "reference": ("referencia material", "refer ncia material"),
    "unit": ("unid med", "un", "unid med "),
    "supplier": ("fornecedor",),
    "country": ("pais fornecedor", "pa s fornecedor", "pais", "pa s"),
}


@dataclass(frozen=True)
class ParsedPrice:
    modality: str
    raw_value: str | None
    amount: Decimal | None
    state: str


@dataclass
class ParsedRow:
    row_number: int
    raw_row: dict[str, Any]
    external_code: str | None
    name: str | None
    manufacturer_code: str | None = None
    reference: str | None = None
    source_url: str | None = None
    brand: str | None = None
    description_raw: str | None = None
    description_display: str | None = None
    application: str | None = None
    category: str | None = None
    subcategory: str | None = None
    unit: str | None = None
    supplier: str | None = None
    country: str | None = None
    availability_raw: str | None = None
    availability_state: str = "nao_informado"
    stock_quantity: Decimal | None = None
    prices: list[ParsedPrice] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    fingerprint: str = ""


@dataclass
class CatalogPreview:
    path: Path
    company_code: str
    file_sha256: str
    file_size: int
    sheet_name: str
    header_row: int
    headers: list[str]
    mapping: dict[str, str]
    rows: list[ParsedRow]

    @property
    def valid_rows(self) -> int:
        return sum(not row.errors for row in self.rows)

    @property
    def rejected_rows(self) -> int:
        return sum(bool(row.errors) for row in self.rows)

    @property
    def warning_count(self) -> int:
        return sum(len(row.warnings) for row in self.rows)

    def summary(self) -> dict[str, Any]:
        return {
            "file": self.path.name,
            "company": self.company_code,
            "sha256": self.file_sha256,
            "size": self.file_size,
            "sheet": self.sheet_name,
            "header_row": self.header_row,
            "headers": self.headers,
            "mapping": self.mapping,
            "total_rows": len(self.rows),
            "valid_rows": self.valid_rows,
            "rejected_rows": self.rejected_rows,
            "warning_count": self.warning_count,
            "sample": [preview_row(row) for row in self.rows[:20]],
        }


@dataclass(frozen=True)
class ImportResult:
    import_id: uuid.UUID
    idempotent: bool
    total_rows: int
    inserted: int
    updated: int
    unchanged: int
    warnings: int
    rejected: int


def canonical_text(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value).strip().casefold())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def cleaned_text(value: Any) -> str | None:
    text = text_value(value)
    if text is None:
        return None
    text = text.replace("_x000D_", "\n")
    # Algumas fontes trazem entidades duplamente escapadas, como &amp;nbsp;.
    for _ in range(2):
        unescaped = html.unescape(text)
        if unescaped == text:
            break
        text = unescaped
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() or None


def text_value(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return str(value)
    return value


def parse_decimal(value: Any) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise InvalidOperation
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            raise InvalidOperation
        return Decimal(str(value))
    text = str(value).strip().replace("R$", "").replace("\u00a0", "").replace(" ", "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(".", "").replace(",", ".")
    return Decimal(text)


def parsed_price(modality: str, value: Any) -> ParsedPrice:
    raw = text_value(value)
    try:
        amount = parse_decimal(value)
    except (InvalidOperation, ValueError):
        return ParsedPrice(modality, raw, None, "invalido")
    if amount is None or amount == 0:
        return ParsedPrice(modality, raw, None, "nao_informado")
    if amount < 0:
        return ParsedPrice(modality, raw, None, "invalido")
    return ParsedPrice(modality, raw, amount, "valido")


def availability_from_text(value: Any) -> tuple[str | None, str]:
    raw = text_value(value)
    normalized = canonical_text(raw)
    if normalized == "em estoque":
        return raw, "em_estoque"
    if normalized == "fora de estoque":
        return raw, "fora_estoque"
    if "backorder" in normalized:
        return raw, "sob_encomenda"
    return raw, "nao_informado"


def availability_from_stock(value: Any) -> tuple[str | None, str, Decimal | None, str | None]:
    raw = text_value(value)
    try:
        quantity = parse_decimal(value)
    except (InvalidOperation, ValueError):
        return raw, "nao_informado", None, "Estoque numérico inválido"
    if quantity is None:
        return raw, "nao_informado", None, None
    return raw, ("em_estoque" if quantity > 0 else "fora_estoque"), quantity, None


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_xlsx(path: Path, settings: Settings) -> None:
    if path.suffix.casefold() != ".xlsx":
        raise CatalogValidationError("Somente XLSX é aceito neste importador inicial.")
    if not path.is_file():
        raise CatalogValidationError("Arquivo não encontrado.")
    if path.stat().st_size > settings.max_upload_mib * 1024 * 1024:
        raise CatalogValidationError("Arquivo excede o limite configurado.")
    try:
        with zipfile.ZipFile(path) as archive:
            total_uncompressed = sum(item.file_size for item in archive.infolist())
            if total_uncompressed > settings.max_upload_mib * 20 * 1024 * 1024:
                raise CatalogValidationError("Conteúdo descompactado excede o limite seguro.")
    except zipfile.BadZipFile as exc:
        raise CatalogValidationError("XLSX inválido.") from exc


def locate_header(workbook) -> tuple[Any, int, list[Any]]:
    for worksheet in workbook.worksheets:
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=1, max_row=10, values_only=True), start=1
        ):
            normalized = {canonical_text(value) for value in values if value is not None}
            is_eco = {"codigo", "nome", "link"}.issubset(normalized)
            has_material = bool({"cod material", "c d material"} & normalized)
            has_sale_price = bool({"preco venda", "pre o venda"} & normalized)
            if is_eco or (has_material and has_sale_price):
                return worksheet, row_number, list(values)
    raise CatalogValidationError("Cabeçalho de catálogo não reconhecido nas primeiras 10 linhas.")


def map_headers(headers: list[Any], company_code: str) -> dict[str, int]:
    by_name: dict[str, int] = {}
    for index, header in enumerate(headers):
        normalized = canonical_text(header)
        if normalized and normalized not in by_name:
            by_name[normalized] = index
    mapping: dict[str, int] = {}
    for field_name, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if alias in by_name:
                mapping[field_name] = by_name[alias]
                break
    required = {"external_code", "name"}
    if company_code == "eco-automacao":
        required |= {"pix_price", "card_price", "availability"}
    else:
        required |= {"sale_price", "stock"}
    missing = sorted(required - mapping.keys())
    if missing:
        raise CatalogValidationError(f"Colunas obrigatórias ausentes: {', '.join(missing)}")
    return mapping


def parse_catalog(path: Path, company_code: str, settings: Settings) -> CatalogPreview:
    if company_code not in {"eco-automacao", "portwell", "choice-tech"}:
        raise CatalogValidationError("Empresa não suportada pelo mapeamento inicial.")
    validate_xlsx(path, settings)
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet, header_row, header_values = locate_header(workbook)
    mapping_indexes = map_headers(header_values, company_code)
    headers = [text_value(value) or f"coluna_{index + 1}" for index, value in enumerate(header_values)]
    mapping = {field_name: headers[index] for field_name, index in mapping_indexes.items()}
    rows: list[ParsedRow] = []
    first_valid_row_by_code: dict[str, int] = {}
    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        if len(rows) >= settings.max_import_rows:
            raise CatalogValidationError("Quantidade de linhas excede o limite configurado.")
        row_values = list(values) + [None] * max(0, len(headers) - len(values))
        raw = {headers[index]: json_value(row_values[index]) for index in range(len(headers))}
        get = lambda field_name: row_values[mapping_indexes[field_name]] if field_name in mapping_indexes else None
        row = ParsedRow(
            row_number=row_number,
            raw_row=raw,
            external_code=text_value(get("external_code")),
            name=text_value(get("name")),
            manufacturer_code=text_value(get("manufacturer_code")),
            reference=text_value(get("reference")),
            source_url=text_value(get("source_url")),
            brand=text_value(get("brand")),
            description_raw=text_value(get("description")),
            application=text_value(get("application")),
            category=text_value(get("category")),
            subcategory=text_value(get("subcategory")),
            unit=text_value(get("unit")),
            supplier=text_value(get("supplier")),
            country=text_value(get("country")),
        )
        if not row.external_code:
            row.errors.append("Código externo ausente")
        if not row.name:
            row.errors.append("Nome ausente")
        if not row.errors and row.external_code:
            first_row_number = first_valid_row_by_code.get(row.external_code)
            if first_row_number is not None:
                row.errors.append(
                    f"Código externo duplicado no arquivo; primeira ocorrência na linha {first_row_number}"
                )
            else:
                first_valid_row_by_code[row.external_code] = row_number
        if not row.brand:
            row.warnings.append("Marca/fabricante ausente")
        if not row.description_raw and not row.application:
            row.warnings.append("Descrição/aplicação ausente")
        row.description_display = cleaned_text(row.description_raw or row.application)
        if company_code == "eco-automacao":
            row.availability_raw, row.availability_state = availability_from_text(get("availability"))
            row.prices = [
                parsed_price("pix", get("pix_price")),
                parsed_price("cartao", get("card_price")),
            ]
        else:
            (
                row.availability_raw,
                row.availability_state,
                row.stock_quantity,
                stock_warning,
            ) = availability_from_stock(get("stock"))
            if stock_warning:
                row.warnings.append(stock_warning)
            row.prices = [parsed_price("venda", get("sale_price"))]
        for price in row.prices:
            if price.state == "invalido":
                row.warnings.append(f"Preço {price.modality} inválido")
        fingerprint_payload = {
            key: value
            for key, value in asdict(row).items()
            if key not in {"raw_row", "warnings", "errors", "fingerprint", "row_number"}
        }
        row.fingerprint = hashlib.sha256(
            json.dumps(fingerprint_payload, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")
        ).hexdigest()
        rows.append(row)
    workbook.close()
    return CatalogPreview(
        path=path,
        company_code=company_code,
        file_sha256=file_sha256(path),
        file_size=path.stat().st_size,
        sheet_name=worksheet.title,
        header_row=header_row,
        headers=headers,
        mapping=mapping,
        rows=rows,
    )


def preview_row(row: ParsedRow) -> dict[str, Any]:
    return {
        "row_number": row.row_number,
        "external_code": row.external_code,
        "name": row.name,
        "brand": row.brand,
        "category": row.category,
        "availability": row.availability_state,
        "prices": [
            {"modality": price.modality, "amount": str(price.amount) if price.amount is not None else None, "state": price.state}
            for price in row.prices
        ],
        "warnings": row.warnings,
        "errors": row.errors,
    }


def normalized_catalog_name(value: str) -> str:
    return canonical_text(value)[:255]


def normalized_brand_name(value: str) -> str:
    # Pontuação em razões sociais não deve criar outra identidade de marca:
    # "S.R.L." e "SRL", por exemplo, precisam consultar a mesma chave de alias.
    return canonical_text(value.replace(".", ""))[:255]


def publish_catalog(
    db: Session,
    preview: CatalogPreview,
    company: Company,
    *,
    currency: str = "BRL",
    observed_at: datetime | None = None,
    created_by_id: uuid.UUID | None = None,
) -> ImportResult:
    existing_import = db.scalar(
        select(CatalogImport).where(
            CatalogImport.company_id == company.id,
            CatalogImport.file_sha256 == preview.file_sha256,
        )
    )
    if existing_import is not None:
        return ImportResult(
            existing_import.id,
            True,
            existing_import.total_rows,
            existing_import.inserted_count,
            existing_import.updated_count,
            existing_import.unchanged_count,
            existing_import.warning_count,
            existing_import.rejected_count,
        )

    import_id = uuid.uuid4()
    batch = CatalogImport(
        id=import_id,
        organization_id=company.organization_id,
        company_id=company.id,
        created_by_id=created_by_id,
        original_filename=preview.path.name,
        source_path=str(preview.path.as_posix()),
        file_sha256=preview.file_sha256,
        file_size=preview.file_size,
        sheet_name=preview.sheet_name,
        header_row=preview.header_row,
        mapping=preview.mapping,
        currency=currency.upper(),
        observed_at=observed_at,
        status="processing",
        total_rows=len(preview.rows),
        inserted_count=0,
        updated_count=0,
        unchanged_count=0,
        warning_count=preview.warning_count,
        rejected_count=preview.rejected_rows,
    )
    db.add(batch)

    offers = {
        offer.external_code: offer
        for offer in db.scalars(select(Offer).where(Offer.company_id == company.id))
    }
    enrichments = {
        entry.offer_id: entry
        for entry in db.scalars(
            select(OfferEnrichment)
            .join(Offer, Offer.id == OfferEnrichment.offer_id)
            .where(Offer.company_id == company.id)
        )
    }
    brand_records = list(
        db.scalars(select(Brand).where(Brand.organization_id == company.organization_id))
    )
    brands = {brand.normalized_name: brand for brand in brand_records}
    brands_by_id = {brand.id: brand for brand in brand_records}
    brand_aliases = {
        alias.normalized_alias: brands_by_id[alias.brand_id]
        for alias in db.scalars(
            select(BrandAlias).where(
                BrandAlias.organization_id == company.organization_id
            )
        )
        if alias.brand_id in brands_by_id
    }
    categories = {
        category.normalized_name: category
        for category in db.scalars(select(Category).where(Category.organization_id == company.organization_id))
    }
    inserted = updated = unchanged = 0

    for row in preview.rows:
        if row.errors:
            db.add(
                ImportRow(
                    id=uuid.uuid4(), catalog_import_id=import_id, sheet_name=preview.sheet_name,
                    row_number=row.row_number, raw_row=row.raw_row, warnings=row.warnings,
                    errors=row.errors, published=False,
                )
            )
            continue

        brand_id = None
        if row.brand:
            brand_key = normalized_brand_name(row.brand)
            brand = brand_aliases.get(brand_key) or brands.get(brand_key)
            if brand is None:
                brand = Brand(
                    id=uuid.uuid4(), organization_id=company.organization_id,
                    canonical_name=row.brand[:255], normalized_name=brand_key,
                )
                brands[brand_key] = brand
                db.add(brand)
                db.add(
                    BrandAlias(
                        id=uuid.uuid4(),
                        organization_id=company.organization_id,
                        brand_id=brand.id,
                        alias=row.brand[:255],
                        normalized_alias=brand_key,
                        source="importado",
                        review_status="sugerido",
                    )
                )
            brand_id = brand.id

        category_id = None
        if row.category:
            category_key = normalized_catalog_name(row.category)
            category = categories.get(category_key)
            if category is None:
                category = Category(
                    id=uuid.uuid4(), organization_id=company.organization_id,
                    name=row.category[:255], normalized_name=category_key,
                    source="importado", review_status="sugerido",
                )
                categories[category_key] = category
                db.add(category)
            category_id = category.id

        offer = offers.get(row.external_code)
        offer_values = {
            "brand_id": brand_id,
            "category_id": category_id,
            "manufacturer_code": row.manufacturer_code,
            "reference": row.reference,
            "name_raw": row.name,
            "description_raw": row.description_raw,
            "description_display": row.description_display,
            "application_raw": row.application,
            "brand_raw": row.brand,
            "category_raw": row.category,
            "subcategory_raw": row.subcategory,
            "source_url": row.source_url,
            "unit_raw": row.unit,
            "supplier_raw": row.supplier,
            "country_raw": row.country,
            "availability_raw": row.availability_raw,
            "availability_state": row.availability_state,
            "stock_quantity": row.stock_quantity,
            "source_fingerprint": row.fingerprint,
        }
        changed = False
        if offer is None:
            offer = Offer(
                id=uuid.uuid4(), organization_id=company.organization_id,
                company_id=company.id, external_code=row.external_code,
                active=True,
                excluded_from_analytics=(company.code == "eco-automacao" and row.external_code == "4574"),
                exclusion_reason=("Registro de teste identificado na fonte" if company.code == "eco-automacao" and row.external_code == "4574" else None),
                **offer_values,
            )
            offers[row.external_code] = offer
            db.add(offer)
            inserted += 1
            changed = True
        elif offer.source_fingerprint != row.fingerprint:
            for field_name, value in offer_values.items():
                setattr(offer, field_name, value)
            offer.active = True
            updated += 1
            changed = True
        else:
            unchanged += 1

        content_sha256 = product_content_sha256(
            name=row.name, manufacturer_code=row.manufacturer_code,
            reference=row.reference, description=row.description_raw,
            application=row.application, brand=row.brand, category=row.category,
        )
        enrichment = enrichments.get(offer.id)
        if enrichment is None or enrichment.content_sha256 != content_sha256:
            treatment = treat_raw(
                name=row.name, manufacturer_code=row.manufacturer_code,
                reference=row.reference, description=row.description_raw,
                application=row.application,
            )
            if enrichment is None:
                enrichment = OfferEnrichment(id=uuid.uuid4(), offer_id=offer.id)
                db.add(enrichment)
                enrichments[offer.id] = enrichment
            enrichment.source = "regra"
            enrichment.source_filename = preview.path.name[:255]
            enrichment.source_sha256 = preview.file_sha256
            enrichment.source_row = row.row_number
            enrichment.content_sha256 = content_sha256
            enrichment.family = treatment.family
            enrichment.family_group = treatment.family_group
            enrichment.model_reference = treatment.model_reference
            enrichment.normalized_title = treatment.normalized_title
            enrichment.search_codes = treatment.search_codes
            enrichment.attributes = treatment.attributes
            enrichment.confidence = treatment.confidence

        db.add(
            ImportRow(
                id=uuid.uuid4(), catalog_import_id=import_id, offer_id=offer.id,
                sheet_name=preview.sheet_name, row_number=row.row_number,
                raw_row=row.raw_row, warnings=row.warnings, errors=[], published=True,
            )
        )
        if changed:
            observation_id = uuid.uuid4()
            db.add(
                OfferObservation(
                    id=observation_id, offer_id=offer.id, catalog_import_id=import_id,
                    source_fingerprint=row.fingerprint, availability_raw=row.availability_raw,
                    availability_state=row.availability_state, stock_quantity=row.stock_quantity,
                    currency=currency.upper(), observed_at=observed_at,
                )
            )
            for price in row.prices:
                db.add(
                    OfferPrice(
                        id=uuid.uuid4(), observation_id=observation_id,
                        modality=price.modality, amount=price.amount,
                        raw_value=price.raw_value, state=price.state,
                    )
                )

    batch.inserted_count = inserted
    batch.updated_count = updated
    batch.unchanged_count = unchanged
    batch.status = "published"
    batch.published_at = datetime.now(timezone.utc)
    return ImportResult(
        import_id, False, len(preview.rows), inserted, updated, unchanged,
        preview.warning_count, preview.rejected_rows,
    )
