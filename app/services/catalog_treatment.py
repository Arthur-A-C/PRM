"""Tratamento determinístico inspirado nos campos e regras do PRM_Pipeline.

Os XLSX tratados fornecem atributos explícitos. Para novas planilhas brutas,
extraímos somente evidências textuais conservadoras antes do pareamento.
"""

import hashlib
import json
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook


def plain(value: str | None) -> str:
    text = unicodedata.normalize("NFKD", value or "").encode("ascii", "ignore").decode()
    return " ".join(re.sub(r"[^a-zA-Z0-9]+", " ", text).lower().split())


FAMILIES = (
    (r"\bbloco de contato\b", "Bloco de Contato", "Comando e Sinalização (botões)"),
    (r"\b(botao de emergencia|botoeira de emergencia)\b", "Botão de Emergência", "Comando e Sinalização (botões)"),
    (r"\b(botao|botoeira|sinaleiro|seletora)\b", "Comando e Sinalização", "Comando e Sinalização (botões)"),
    (r"\b(cortina de luz|rele de seguranca|sensor de seguranca)\b", "Segurança de Máquinas", "Segurança de Máquinas"),
    (r"\b(sensor|transdutor)\b", "Sensor", "Sensores"),
    (r"\bencoder\b", "Encoder", "Encoders"),
    (r"\b(fonte de alimentacao|nobreak|ups)\b", "Fonte de Alimentação", "Fontes e UPS"),
    (r"\b(clp|ihm|gateway)\b", "Automação e CLP", "Automação e CLP"),
    (r"\b(rele|relé)\b", "Relé", "Relés"),
    (r"\b(disjuntor|fusivel|contator)\b", "Proteção Elétrica", "Proteção Elétrica"),
    (r"\b(cabo|conector|borne)\b", "Cabo ou Conector", "Cabos e Conectores"),
)

ATTR_ALIASES = {
    "grau de protecao ip": "protecao_ip",
    "grau de protecao": "protecao_ip",
    "grau ip": "protecao_ip",
    "ip": "protecao_ip",
    "protocolos de comunicacao": "protocolo_de_comunicacao",
    "protocolo": "protocolo_de_comunicacao",
    "diametro": "diametro",
    "posicoes": "posicoes",
    "distancia maxima": "distancia_maxima",
    "distancia max": "distancia_maxima",
    "tensao de entrada": "tensao_de_entrada",
    "tensao de saida": "tensao_de_saida",
}


def attribute_key(label: str) -> str:
    raw = label[4:] if label.startswith("ATR_") else label
    normalized = plain(raw.replace("_", " "))
    return ATTR_ALIASES.get(normalized, normalized.replace(" ", "_")[:100])


def product_content_sha256(
    *, name: str | None, manufacturer_code: str | None, reference: str | None,
    description: str | None, application: str | None, brand: str | None,
    category: str | None,
) -> str:
    fields = (name, manufacturer_code, reference, description, application, brand, category)
    return hashlib.sha256(json.dumps(fields, ensure_ascii=False).encode("utf-8")).hexdigest()


@dataclass(frozen=True)
class Treatment:
    family: str | None = None
    family_group: str | None = None
    model_reference: str | None = None
    normalized_title: str | None = None
    search_codes: list[str] = field(default_factory=list)
    attributes: dict[str, str] = field(default_factory=dict)
    confidence: Decimal | None = None


def treat_raw(
    *, name: str, manufacturer_code: str | None, reference: str | None,
    description: str | None, application: str | None,
) -> Treatment:
    name_normalized = plain(name)
    family = group = None
    for pattern, label, group_name in FAMILIES:
        if re.search(pattern, name_normalized[:160]):
            family, group = label, group_name
            break
    codes = []
    for value in (manufacturer_code, reference):
        if value and plain(value) not in {"nulo", "nao deletar", "nao informado"}:
            value = value.strip()
            if value and value not in codes:
                codes.append(value[:255])
    evidence = " ".join(filter(None, (name, description, application)))[:4000]
    attributes = {}
    ip = re.search(r"\bIP\s*([0-9]{2,3})\b", evidence, re.I)
    if ip:
        attributes["protecao_ip"] = "IP" + ip.group(1)
    voltage = re.search(r"(?<![\w.,])(\d{1,3}(?:[.,]\d+)?)\s*V\s*(CC|CA|DC|AC)?\b", evidence, re.I)
    if voltage:
        kind = {"CC": "Vcc", "DC": "Vcc", "CA": "Vca", "AC": "Vca"}.get((voltage.group(2) or "").upper(), "V")
        attributes["tensao"] = voltage.group(1).replace(",", ".") + " " + kind
    return Treatment(
        family=family, family_group=group,
        model_reference=codes[0] if codes else None,
        normalized_title=name_normalized or None,
        search_codes=codes,
        attributes=attributes,
        confidence=Decimal("0.60") if codes or attributes or family else Decimal("0.30"),
    )


@dataclass(frozen=True)
class TreatedRow:
    code: str
    row_number: int
    treatment: Treatment


def _parse_attributes(value: object) -> dict[str, str]:
    result = {}
    for part in str(value or "").split(" | "):
        if "=" not in part:
            continue
        name, raw_value = part.split("=", 1)
        key = attribute_key(name)
        if key and raw_value.strip():
            result[key] = raw_value.strip()[:500]
    return result


def iter_treated_xlsx(path: Path, company_code: str, *, max_rows: int = 50000) -> Iterator[TreatedRow]:
    if path.suffix.lower() != ".xlsx" or not path.is_file():
        raise ValueError("É necessário um arquivo XLSX existente.")
    with zipfile.ZipFile(path) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > 400 * 1024 * 1024:
            raise ValueError("Conteúdo XLSX descompactado excede o limite.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        expected_sheet = "Base_ECO" if company_code == "eco-automacao" else "Base_PRM"
        if sheet.title != expected_sheet:
            raise ValueError(f"Aba esperada: {expected_sheet}.")
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows)]
        columns = {name: index for index, name in enumerate(headers)}
        code_column = "Codigo_ECO" if company_code == "eco-automacao" else "Cód. Material"
        required = {code_column, "Familia_Produto", "Modelo_Referencia", "Atributos_PRM"}
        if not required.issubset(columns):
            raise ValueError("Planilha tratada sem colunas PRM obrigatórias.")
        attr_columns = [(index, attribute_key(header)) for index, header in enumerate(headers) if header.startswith("ATR_")]
        def get(values, column):
            index = columns.get(column)
            return values[index] if index is not None and index < len(values) else None
        for row_number, values in enumerate(rows, start=2):
            if row_number - 1 > max_rows:
                raise ValueError("Planilha tratada excede o limite de linhas.")
            code = str(get(values, code_column) or "").strip()
            if not code:
                continue
            attributes = _parse_attributes(get(values, "Atributos_PRM"))
            for index, key in attr_columns:
                if key and index < len(values) and values[index] not in (None, ""):
                    attributes[key] = str(values[index]).strip()[:500]
            attributes.update({k: v for k, v in _parse_attributes(get(values, "Outros_Atributos")).items() if k not in attributes})
            codes = [value.strip()[:255] for value in str(get(values, "Codigos_Busca") or "").split("|") if value.strip()]
            model = str(get(values, "Modelo_Referencia") or "").strip()[:255] or None
            if model and model not in codes:
                codes.insert(0, model)
            confidence = None
            try:
                raw_confidence = get(values, "Confianca_Tratamento")
                if raw_confidence not in (None, ""):
                    parsed = Decimal(str(raw_confidence))
                    if Decimal("0") <= parsed <= Decimal("1"):
                        confidence = parsed
            except InvalidOperation:
                pass
            def optional(column):
                value = get(values, column)
                return str(value).strip()[:255] if value not in (None, "") else None
            yield TreatedRow(
                code=code, row_number=row_number,
                treatment=Treatment(
                    family=optional("Familia_Produto"),
                    family_group=optional("Grupo_Familia_PRM"),
                    model_reference=model,
                    normalized_title=optional("Titulo_Normalizado_PRM"),
                    search_codes=list(dict.fromkeys(codes)),
                    attributes=attributes,
                    confidence=confidence,
                ),
            )
    finally:
        workbook.close()
