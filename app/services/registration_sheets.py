"""Review-first XLSX/CSV intake for the current registration-sheet layout."""

from __future__ import annotations

import csv
import hashlib
import io
from collections import Counter

from openpyxl import load_workbook

from app.services.registration_core import fold, from_sheet_row


REQUIRED = {"nome do produto", "titulo", "fornecedor", "codigo cigam", "status"}


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _sheet_number(name: str, fallback: int) -> int:
    import re
    found = re.search(r"\d+", name)
    return int(found.group()) if found else fallback


def _read_xlsx(content: bytes):
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        for ordinal, sheet in enumerate(book.worksheets, 1):
            rows = sheet.iter_rows(values_only=True)
            first = next(rows, ())
            second = next(rows, ())
            headers = next(rows, ())
            if not any("Voltar ao" in _text(cell) for cell in first):
                continue
            category = next((_text(cell) for cell in second if _text(cell)), "")
            yield sheet.title, ordinal, category, [_text(cell) for cell in headers], (
                tuple(_text(cell) for cell in row) for row in rows
            )
    finally:
        book.close()


def _read_csv(content: bytes):
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        decoded = content.decode("cp1252")
    reader = csv.reader(io.StringIO(decoded, newline=""))
    first = next(reader, ())
    second = next(reader, ())
    headers = next(reader, ())
    if not any("Voltar ao" in cell for cell in first):
        raise ValueError("CSV não segue o layout de cadastro com índice na linha 1.")
    category = next((cell.strip() for cell in second if cell.strip()), "")
    yield "CSV", 1, category, headers, reader


def preview_sheet(
    content: bytes, filename: str, company: str,
    existing_skus: set[str] | None = None,
    store_names: dict[str, str] | None = None,
) -> dict:
    if not content or len(content) > 20 * 1024 * 1024:
        raise ValueError("Planilha vazia ou acima de 20 MiB.")
    extension = filename.lower().rsplit(".", 1)[-1]
    if extension not in {"xlsx", "csv"}:
        raise ValueError("Envie um XLSX ou CSV.")
    source = _read_xlsx(content) if extension == "xlsx" else _read_csv(content)
    seen_tabs: set[str] = set()
    used, duplicate_tabs, conflicts = [], [], []
    grouped: dict[str, list[tuple[int, int, dict]]] = {}
    line_count = 0
    repeated = Counter()
    for tab_name, ordinal, category, headers, rows in source:
        normalized_headers = [fold(header) for header in headers]
        if not REQUIRED.issubset(normalized_headers):
            conflicts.append(f"Aba {tab_name}: cabeçalho obrigatório ausente.")
            continue
        if len(headers) != len(set(normalized_headers)):
            conflicts.append(f"Aba {tab_name}: cabeçalho duplicado.")
            continue
        buffered = []
        for index, values in enumerate(rows, 4):
            if not any(_text(value) for value in values):
                continue
            line_count += 1
            row = dict(zip(headers, (_text(value) for value in values)))
            buffered.append((index, row))
        fingerprint = hashlib.sha256(repr((category, headers, buffered)).encode("utf-8")).hexdigest()
        if fingerprint in seen_tabs:
            duplicate_tabs.append(tab_name)
            continue
        seen_tabs.add(fingerprint)
        used.append(tab_name)
        for row_number, row in buffered:
            sku = next((value for name, value in row.items() if fold(name) == "codigo cigam"), "")
            if not sku:
                conflicts.append(f"{tab_name}:{row_number}: Código CIGAM ausente.")
                continue
            if "e+" in sku.casefold():
                conflicts.append(f"{tab_name}:{row_number}: SKU em notação científica requer fonte textual.")
                continue
            grouped.setdefault(sku, []).append((ordinal, row_number, {"row": row, "category": category, "tab": tab_name}))
            repeated[sku] += 1
    products = []
    for sku, occurrences in grouped.items():
        def priority(item):
            row = item[2]["row"]
            status = next((value for name, value in row.items() if fold(name) == "status"), "")
            return (fold(status) != "concluido", _sheet_number(item[2]["tab"], item[0]), item[1])
        ordered = sorted(occurrences, key=priority)
        chosen = ordered[0][2]
        titles = {
            next((value for name, value in item[2]["row"].items() if fold(name) == "titulo"), "")
            for item in ordered
        }
        if len(titles) > 1:
            conflicts.append(
                f"SKU {sku}: títulos divergentes; priorizado {chosen['tab']}:{ordered[0][1]} "
                "(status Concluído, depois menor número de aba)."
            )
        try:
            product = from_sheet_row(company, chosen["row"], chosen["category"], store_names=store_names)
        except ValueError as exc:
            conflicts.append(f"SKU {sku}: {exc}")
            continue
        for _, _, occurrence in ordered[1:]:
            try:
                other = from_sheet_row(company, occurrence["row"], occurrence["category"], store_names=store_names)
            except ValueError as exc:
                conflicts.append(f"SKU {sku}, aba {occurrence['tab']}: {exc}")
                continue
            for category in other.category_paths:
                if category not in product.category_paths:
                    product.category_paths.append(category)
            by_name = {fold(item.name): item for item in product.attributes}
            for item in other.attributes:
                key = fold(item.name)
                if key in by_name:
                    for value in item.values:
                        if value not in by_name[key].values:
                            by_name[key].values.append(value)
                elif len(product.attributes) < 15:
                    product.attributes.append(item)
                    by_name[key] = item
                else:
                    conflicts.append(f"SKU {sku}: atributo excedente após consolidação: {item.name}.")
        products.append(product)
    known = existing_skus or set()
    return {
        "total_rows": line_count, "used_tabs": used, "duplicate_tabs": duplicate_tabs,
        "distinct_products": len(products),
        "repeated_skus": {sku: count for sku, count in repeated.items() if count > 1},
        "existing_count": sum(product.sku in known for product in products) if existing_skus is not None else None,
        "new_count": sum(product.sku not in known for product in products) if existing_skus is not None else None,
        "conflicts": conflicts,
        "products": [product.model_dump() for product in products],
        "examples": [product.model_dump() for product in products[:3]],
    }
