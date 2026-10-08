"""Auditoria streaming dos seis arquivos de origem/sidecar, sem escrita no banco."""

from __future__ import annotations

from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.services.catalog_import import file_sha256, text_value


@dataclass(frozen=True)
class AuditSource:
    label: str
    relative_path: str
    company: str
    layer: str
    sheet: str
    header_row: int
    code_column: int  # zero-based; checked against the workbook header
    code_header: str


SOURCES = (
    AuditSource("choice_raw", "produtos-choice.xlsx", "choice-tech", "raw", "Sheet1", 3, 4, "Material"),
    AuditSource("choice_treated", "treated/Produtos-Choice.xlsx", "choice-tech", "treated", "Base_PRM", 1, 0, "Material"),
    AuditSource("portwell_raw", "produtos-portwell.xlsx", "portwell", "raw", "Export", 1, 4, "Material"),
    AuditSource("portwell_treated", "treated/Produtos-Portwell.xlsx", "portwell", "treated", "Base_PRM", 1, 0, "Material"),
    AuditSource("eco_raw", "produtos_ecoautomacao.xlsx", "eco-automacao", "raw", "Produtos", 1, 0, "Codigo"),
    AuditSource("eco_treated", "treated/Produtos-ECO.xlsx", "eco-automacao", "treated", "Base_ECO", 1, 0, "Codigo"),
)

CASE_CODES = {
    "choice-tech": {"111001000409", "190003000491", "112005000072", "112008000051"},
    "portwell": {"316001000891", "316001000892", "311006000015", "312001000032", "317010000010"},
    "eco-automacao": {"4574", "P024185", "P024190", "P023972", "P000194", "P002212", "P002288"},
}


def _source_audit(root: Path, source: AuditSource) -> tuple[dict[str, Any], Counter[str]]:
    path = root / source.relative_path
    if not path.is_file():
        raise FileNotFoundError(f"Arquivo de auditoria ausente: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if source.sheet not in workbook.sheetnames:
            raise ValueError(f"{path.name}: aba {source.sheet!r} ausente")
        sheet = workbook[source.sheet]
        sheet.reset_dimensions()  # A exportação Portwell não informa dimensões.
        physical = missing_code = blank = 0
        first_line: int | None = None
        last_line: int | None = None
        headers: list[str] = []
        codes: Counter[str] = Counter()
        cases: dict[str, list[int]] = {}
        non_product: list[dict[str, Any]] = []
        commercial_first: dict[str, tuple[int, tuple[str | None, ...]]] = {}
        commercial_conflicts: dict[str, dict[str, Any]] = {}
        zero_prices: Counter[str] = Counter()
        availability: Counter[str] = Counter()
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            if row_number == source.header_row:
                headers = [text_value(value) or "" for value in row]
                if source.code_column >= len(headers) or source.code_header.casefold() not in headers[source.code_column].casefold():
                    raise ValueError(f"{path.name}: coluna de código inesperada no cabeçalho")
                continue
            if row_number < source.header_row:
                continue
            if not any(value is not None and str(value).strip() for value in row):
                blank += 1
                continue
            physical += 1
            first_line = row_number if first_line is None else first_line
            last_line = row_number
            raw_code = row[source.code_column] if source.code_column < len(row) else None
            code = text_value(raw_code)
            if code is None:
                missing_code += 1
                non_product.append({"line": row_number, "first_cell": text_value(row[0]) if row else None})
                continue
            codes[code] += 1
            if code in CASE_CODES[source.company]:
                cases.setdefault(code, []).append(row_number)
            if source.label == "portwell_raw":
                fields = ("indice_mk", "indice_fator", "preco_compra", "preco_venda")
                current = tuple(text_value(row[index]) if index < len(row) else None for index in (22, 23, 24, 25))
                if code in commercial_first:
                    first_line, first_values = commercial_first[code]
                    differences = [field for field, left, right in zip(fields, first_values, current) if left != right]
                    commercial_conflicts[code] = {"lines": [first_line, row_number], "different_fields": differences}
                else:
                    commercial_first[code] = (row_number, current)
            if source.label == "eco_raw":
                for modality, index in (("pix", 4), ("cartao", 5)):
                    if index < len(row) and row[index] == 0:
                        zero_prices[modality] += 1
                availability[text_value(row[6]) or ""] += 1
        report = {
            **asdict(source),
            "sha256": file_sha256(path),
            "size_bytes": path.stat().st_size,
            "columns": len(headers),
            "physical_rows": physical,
            "coded_rows": sum(codes.values()),
            "distinct_codes": len(codes),
            "blank_rows_after_header": blank,
            "missing_code_rows": missing_code,
            "non_product_rows": non_product,
            "duplicate_codes": {code: count for code, count in codes.items() if count > 1},
            "commercial_conflicts": commercial_conflicts,
            "zero_prices": dict(zero_prices),
            "availability_counts": dict(availability),
            "case_lines": cases,
            "first_data_line": first_line,
            "last_data_line": last_line,
        }
        return report, codes
    finally:
        workbook.close()


def audit_six_files(root: Path) -> dict[str, Any]:
    """Reconcilia sidecars por multiplicidade do código, nunca por posição."""
    results: dict[str, dict[str, Any]] = {}
    counters: dict[str, Counter[str]] = {}
    for source in SOURCES:
        results[source.label], counters[source.label] = _source_audit(root, source)
    reconciliation = {}
    for company, prefix in (("choice-tech", "choice"), ("portwell", "portwell"), ("eco-automacao", "eco")):
        raw, treated = counters[f"{prefix}_raw"], counters[f"{prefix}_treated"]
        reconciliation[company] = {
            "raw_only": dict(raw - treated),
            "treated_only": dict(treated - raw),
            "same_code_multiplicity": raw == treated,
        }
    return {
        "audit_version": "1",
        "mode": "read_only_staging",
        "sources": results,
        "reconciliation": reconciliation,
        "product_rows_per_layer": sum(results[f"{prefix}_raw"]["coded_rows"] for prefix in ("choice", "portwell", "eco")),
    }
