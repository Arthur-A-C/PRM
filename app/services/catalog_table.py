"""Mapped tabular input for the existing catalog publisher (no database writes)."""
import csv
import hashlib
import json
import zipfile
from dataclasses import asdict
from pathlib import Path

from openpyxl import load_workbook

from app.services.catalog_import import (
    CatalogPreview, ParsedRow, FIELD_ALIASES, canonical_text, text_value,
    cleaned_text, parsed_price, availability_from_text, availability_from_stock, file_sha256,
)

FIELDS = {"external_code": "SKU / código", "name": "Nome", "brand": "Marca",
          "category": "Família / categoria", "subcategory": "Segmento",
          "description": "Descrição", "application": "Aplicação",
          "sale_price": "Preço venda", "pix_price": "Preço Pix", "card_price": "Preço cartão",
          "availability": "Disponibilidade", "stock": "Quantidade em estoque",
          "manufacturer_code": "Código fabricante", "reference": "Referência",
          "source_url": "URL", "unit": "Unidade", "supplier": "Fornecedor", "country": "País"}


def read_table(path, *, sheet="", max_rows=50000, header_row=1):
    """Read one sheet, never evaluate formulas/macros; bound expanded XLSX content."""
    suffix = path.suffix.lower()
    rows, sheets = [], []
    if not 1 <= header_row <= 100:
        raise ValueError("Linha de cabeçalho deve estar entre 1 e 100.")
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            sample = stream.read(8192); stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            for index, row in enumerate(csv.reader(stream, dialect)):
                if index >= max_rows + header_row:
                    raise ValueError("Limite de linhas excedido.")
                rows.append(row)
        sheet = "CSV"; sheets = [sheet]
    elif suffix == ".xlsx":
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 200 * 1024 * 1024:
                raise ValueError("Excel descompactado excede 200 MiB.")
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            sheets = workbook.sheetnames
            if not sheet and len(sheets) > 1:
                return {"sheets": sheets, "sheet": "", "headers": [], "rows": []}
            sheet = sheet or sheets[0]
            if sheet not in sheets:
                raise ValueError("Aba inválida.")
            if (workbook[sheet].max_column or 0) > 200:
                raise ValueError("Use até 200 colunas.")
            for index, row in enumerate(workbook[sheet].iter_rows(values_only=True)):
                if index >= max_rows + header_row:
                    raise ValueError("Limite de linhas excedido.")
                rows.append([text_value(value) or "" for value in row])
        finally:
            workbook.close()
    elif suffix == ".xls":
        import xlrd
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        try:
            sheets = workbook.sheet_names()
            if not sheet and len(sheets) > 1:
                return {"sheets": sheets, "sheet": "", "headers": [], "rows": []}
            sheet = sheet or sheets[0]
            tab = workbook.sheet_by_name(sheet)
            if tab.ncols > 200:
                raise ValueError("Use até 200 colunas.")
            if tab.nrows > max_rows + header_row:
                raise ValueError("Limite de linhas excedido.")
            rows = [[text_value(v) or "" for v in tab.row_values(i)] for i in range(tab.nrows)]
        finally:
            workbook.release_resources()
    else:
        raise ValueError("Envie CSV, XLSX ou XLS.")
    if len(rows) < header_row or not rows[header_row - 1]:
        raise ValueError("Arquivo sem cabeçalho.")
    headers = [str(v).strip() or f"Coluna {i+1}" for i, v in enumerate(rows[header_row - 1])]
    if len(headers) > 200 or len(headers) != len(set(headers)):
        raise ValueError("Use até 200 colunas, com cabeçalhos distintos.")
    return {"sheets": sheets, "sheet": sheet, "headers": headers,
            "rows": rows[header_row:]}


def suggest_mapping(headers):
    extras = {"external_code": ("sku", "cigam", "codigo externo"), "name": ("produto", "nome produto"),
              "sale_price": ("preco", "valor"), "category": ("familia", "categoria"),
              "subcategory": ("segmento",), "availability": ("disponibilidade",)}
    return {key: header for key in FIELDS for header in headers
            if canonical_text(header) in {canonical_text(v) for v in (*FIELD_ALIASES.get(key, ()), *extras.get(key, ()))}}


def mapped_preview(path, company_code, table, mapping, header_row=1):
    if not {"external_code", "name"} <= mapping.keys():
        raise ValueError("Mapeie SKU / código e Nome.")
    if any(not isinstance(k, str) or not isinstance(v, str) or k not in FIELDS or v not in table["headers"] for k, v in mapping.items()):
        raise ValueError("Mapeamento inválido.")
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("Cada coluna deve ser mapeada uma única vez.")
    rows, seen = [], set()
    for number, values in enumerate(table["rows"], header_row + 1):
        if not any(str(v).strip() for v in values):
            continue
        raw = dict(zip(table["headers"], values))
        get = lambda key: text_value(raw.get(mapping.get(key)))
        row = ParsedRow(number, raw, get("external_code"), get("name"))
        for key in ("brand", "category", "subcategory", "application", "manufacturer_code", "reference", "source_url", "unit", "supplier", "country"):
            setattr(row, key, get(key))
        for key, maximum in {"brand": 500, "category": 500, "subcategory": 500,
                             "manufacturer_code": 255, "reference": 500, "unit": 100,
                             "supplier": 500, "country": 255}.items():
            if len(getattr(row, key) or "") > maximum:
                row.errors.append(f"{FIELDS[key]} acima de {maximum} caracteres")
        row.description_raw = get("description")
        row.description_display = cleaned_text(row.description_raw or row.application)
        if not row.external_code or len(row.external_code) > 255:
            row.errors.append("SKU ausente ou acima de 255 caracteres")
        if not row.name:
            row.errors.append("Nome ausente")
        if row.external_code in seen:
            row.errors.append("SKU duplicado neste arquivo")
        seen.add(row.external_code)
        row.availability_raw, row.availability_state = availability_from_text(get("availability"))
        if "stock" in mapping:
            raw_stock, state, row.stock_quantity, warning = availability_from_stock(get("stock"))
            if not row.availability_raw:
                row.availability_raw, row.availability_state = raw_stock, state
            if warning:
                row.warnings.append(warning)
        row.prices = [parsed_price(modality, get(field)) for field, modality in
                      (("sale_price", "venda"), ("pix_price", "pix"), ("card_price", "cartao")) if field in mapping]
        for price in row.prices:
            if price.state == "invalido":
                row.errors.append(f"Preço {price.modality} inválido")
        content = {k: v for k, v in asdict(row).items() if k not in {"row_number", "raw_row", "errors", "warnings", "fingerprint"}}
        row.fingerprint = hashlib.sha256(json.dumps(content, sort_keys=True, default=str).encode()).hexdigest()
        rows.append(row)
    return CatalogPreview(Path(path), company_code, file_sha256(path), path.stat().st_size,
                          table["sheet"], header_row, table["headers"], mapping, rows)
