"""Durable review/approve/confirm queue for product files (images, WCDA); no browser automation."""
import csv
import hashlib
import io
import json
import shutil
import sqlite3
import stat
import time
import uuid
import zipfile
from pathlib import Path
from contextlib import contextmanager

from openpyxl import load_workbook
from app.services.store_config import ROOT, COMPANIES, StoreSettings
from app.services.store_publish import MIMES, PublishSettings, StagingPublisher, file_info, snapshot_hash
from app.services.store_registration_plan import (prepare_files, file_plan, stale_supplier_datasheet, store_limit_check,
                                                   normalize_scope, main_options, scoped_images, scoped_file_plan,
                                                   classify_by_index, rule_images, folder_inventory)
from app.services.registration_files import find_product_folder
from app.services.registration_woo_api import WooClient, credentials_for
from app.services.datasheet_batches import relative_name, identify_product_folders
from app.services.digitalization_storage import validate_sku
from app.services.store_documents import interpret_documents
from app.services import store_publish_grant as grants
from app.services import store_registration_authorization as authorization
from app.services.store_publish import _protected_hash


def parse_skus(raw, name, *, details=False, strict=False):
    if len(raw)>5*1024**2:
        raise ValueError("Planilha acima de 5 MiB.")
    if name.lower().endswith(".csv"):
        try: text=raw.decode("utf-8-sig")
        except UnicodeError: text=raw.decode("cp1252")
        try: dialect=csv.Sniffer().sniff(text[:2000],delimiters=",;\t")
        except csv.Error: dialect=csv.excel
        rows=list(csv.reader(io.StringIO(text),dialect))
    elif name.lower().endswith((".xlsx",".xlsm")):
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist())>30*1024**2:
                raise ValueError("Planilha descompactada excessiva.")
        book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        try:
            rows=[]
            for row in book.active.iter_rows(values_only=True):
                if len(rows)>1000: raise ValueError("Limite de 1000 linhas.")
                rows.append(["" if v is None else str(int(v)) if isinstance(v,float) and v.is_integer() else str(v) for v in row])
        finally: book.close()
    else: raise ValueError("Envie CSV ou XLSX/XLSM.")
    rows=[r for r in rows if any(str(v).strip() for v in r)]
    if not rows or len(rows)>1000: raise ValueError("Planilha vazia ou acima de 1000 linhas.")
    headers=[v.strip().lower() for v in rows[0]]
    aliases={"sku","código","codigo","cod","cod.","código cigam","codigo cigam"}
    index=next((i for i,h in enumerate(headers) if h in aliases),None)
    if index is None:
        if any(len(r)!=1 for r in rows): raise ValueError("Informe uma coluna SKU/CIGAM.")
        index=0
    else: rows=rows[1:]
    skus=[]; records={}
    for row in rows:
        if index>=len(row) or not row[index].strip(): continue
        sku=validate_sku(row[index] if strict else row[index].strip())
        if "e+" in sku.lower(): raise ValueError("SKU em notação científica; envie como texto.")
        if strict and sku in records: raise ValueError("SKU duplicado na planilha complementar.")
        if sku not in skus: skus.append(sku)
        record={h:row[i] for i,h in enumerate(headers) if i<len(row) and h in
                {"nome","descricao","descricao_curta","categoria","empresa","observacoes"}}
        if sku in records and records[sku]!=record: raise ValueError("Linhas conflitantes para o mesmo SKU.")
        records[sku]=record
    if not skus: raise ValueError("Nenhum SKU.")
    return records if details else skus


BATCH_REPORTS={"relatorio_processamento.csv","erros.csv"}


def batch_bundle(entries, directories, company):
    """Recognize the Geração em Lote export: COMPANY/SKU.pdf, COMPANY/SKU/analysis.json, COMPANY/SKU/fontes/.

    Returns entries renamed to one folder per SKU, or None for any other layout.
    """
    names=[n for n,_ in entries]
    reports={n for n in names if "/" not in n and n.lower().endswith(".csv")}
    tops={n.split("/")[0] for n in names if "/" in n}|{d.split("/")[0] for d in directories}
    if not reports or not reports<=BATCH_REPORTS or len(tops)!=1 or next(iter(tops)) not in {c.upper() for c in COMPANIES}:
        return None
    if next(iter(tops))!=company.upper(): raise ValueError("Pacote do lote pertence a outra empresa.")
    renamed=[]
    for name,info in entries:
        if name in reports: continue
        parts=name.split("/")
        if len(parts)==2 and parts[1].lower().endswith(".pdf"):
            sku=validate_sku(parts[1][:-4]);name=f"{parts[0]}/{sku}/{sku}_Datasheet.pdf"
        renamed.append((name,info))
    if len({n.casefold() for n,_ in renamed})!=len(renamed): raise ValueError("Caminho duplicado.")
    return renamed


def extract_files(raw, target, *, identification="csv", archive_name="", root="", company="", limit_mib=1024, package_mib=2048):
    from app.services.upload_limits import check_package, check_size, copy_bounded, MIB
    stream=io.BytesIO(raw) if isinstance(raw,bytes) else raw
    stream.seek(0,2); check_package(stream.tell(),package_mib);stream.seek(0)
    with zipfile.ZipFile(stream) as archive:
        infos=archive.infolist()
        if len(infos)>3000:
            raise ValueError("ZIP excede o limite de 3000 entradas.")
        check_package(sum(i.file_size for i in infos),package_mib)  # uncompressed total: ZIP-bomb guard
        if shutil.disk_usage(target if target.exists() else target.parent).free < sum(i.file_size for i in infos)+64*MIB:
            raise ValueError("Espaço em disco insuficiente para extrair o ZIP.")
        names=set(); entries=[]; directories=[]
        for info in infos:
            name=relative_name(info.orig_filename)
            if name.casefold() in names: raise ValueError("Caminho duplicado.")
            names.add(name.casefold())
            mode=info.external_attr>>16
            if stat.S_ISLNK(mode) or stat.S_IFMT(mode) not in {0,stat.S_IFREG,stat.S_IFDIR} or info.flag_bits&1:
                raise ValueError("Link ou criptografia não permitidos.")
            if info.is_dir():
                directories.append(name)
                continue
            check_size(info.file_size,limit_mib)
            if info.file_size/max(1,info.compress_size)>200:
                raise ValueError("Arquivo excessivo ou compressão suspeita.")
            suffix=Path(name).suffix.lower()
            if suffix not in {".png",".jpg",".jpeg",".webp",".pdf",".zip",".step",".stp",".dwg",".dxf",".rar",".7z",".svg",".gif",".json",".csv"}:
                raise ValueError("Extensão não permitida.")
            if suffix==".json" and Path(name).name not in {"manifest.json","datasheet_dados.json","analysis.json","product.json"}:
                raise ValueError("JSON não reconhecido; envie apenas dados/manifesto do produto.")
            if suffix in {".json",".csv"} and info.file_size>5*1024**2:
                raise ValueError("Metadados/planilha acima de 5 MiB.")
            entries.append((name,info))
        bundle=batch_bundle(entries,directories,company) if identification=="folder" and not root else None
        if bundle is not None:
            entries=bundle;root=entries[0][0].split("/")[0] if entries else ""
            directories=[d for d in directories if d.split("/")[0]==root]
        prefixes=identify_product_folders([n for n,_ in entries],directories,root,archive_name=archive_name) if identification=="folder" else {}
        if identification=="folder":
            if not prefixes or len(prefixes)>1000: raise ValueError("Nenhum SKU identificado ou limite de 1000 excedido.")
            if len({s.casefold() for s in prefixes})!=len(prefixes): raise ValueError("SKUs duplicados no ZIP.")
            for sku in prefixes:
                validate_sku(sku)
                (target/sku).mkdir(parents=True,exist_ok=True)
        csvs=[(n,i) for n,i in entries if n.lower().endswith('.csv')]
        if len(csvs)>1: raise ValueError("Envie no máximo um CSV complementar no ZIP.")
        for name,info in entries:
            if name.lower().endswith('.csv'): continue
            destination=name
            if prefixes:
                matches=[s for s,prefix in prefixes.items() if name.startswith(prefix)]
                if len(matches)!=1: raise ValueError("Arquivo sem pasta de SKU inequívoca.")
                sku=matches[0];destination=sku+"/"+name[len(prefixes[sku]):]
            path=(target/destination).resolve()
            if not path.is_relative_to(target.resolve()): raise ValueError("Caminho inválido.")
            path.parent.mkdir(parents=True,exist_ok=True)
            with archive.open(info) as reader, path.open('wb') as writer:
                copy_bounded(reader,writer,min(info.file_size,limit_mib*MIB))
        return list(prefixes), (archive.read(csvs[0][1]),csvs[0][0]) if csvs else None


class Jobs:
    def __init__(self, policy=None):
        self.policy=policy or PublishSettings()
        self.root=Path(self.policy.storage_dir)
        if not self.root.is_absolute(): self.root=ROOT/self.root
        self.root=self.root.resolve();self.root.mkdir(parents=True,exist_ok=True)
        with self.db() as db:
            db.executescript("""
            CREATE TABLE IF NOT EXISTS batches(id TEXT PRIMARY KEY,org TEXT,owner TEXT,company TEXT,origin TEXT,created REAL);
            CREATE TABLE IF NOT EXISTS jobs(batch TEXT,sku TEXT,state TEXT,error TEXT DEFAULT '',plan TEXT DEFAULT '{}',
              token TEXT DEFAULT '',lease REAL DEFAULT 0,PRIMARY KEY(batch,sku));
            CREATE TABLE IF NOT EXISTS journal(scope TEXT,key TEXT,value TEXT,PRIMARY KEY(scope,key));
            CREATE TABLE IF NOT EXISTS audit(batch TEXT,sku TEXT,at REAL,actor TEXT,event TEXT,detail TEXT);
            """+grants.SCHEMA)
            columns={r[1] for r in db.execute("PRAGMA table_info(grants)")}
            if "scope" not in columns:
                db.execute("ALTER TABLE grants ADD COLUMN scope TEXT DEFAULT ''")
            # Production removal grants (env) and the approver they were issued against (four-eyes rule).
            if "env" not in columns:
                db.execute("ALTER TABLE grants ADD COLUMN env TEXT DEFAULT 'homologation'")
            if "approver" not in columns:
                db.execute("ALTER TABLE grants ADD COLUMN approver TEXT DEFAULT ''")
            from app.services.store_removal_production import ENV_SCHEMA
            db.executescript(ENV_SCHEMA)

    @contextmanager
    def db(self):
        db=sqlite3.connect(self.root/"jobs.sqlite3",timeout=20)
        db.row_factory=sqlite3.Row
        try:
            with db:
                yield db
        finally:
            db.close()

    def create(self, company, org, owner, skus, archive, settings, *, identification="csv", archive_name="", root="", complements=None, scope=None):
        """scope: what to register ({"gallery","documents","main"}). None keeps the previous plan (Geração em Lote hand-off)."""
        if company not in COMPANIES: raise ValueError("Empresa inválida.")
        scope=normalize_scope(scope) if scope is not None else None
        if identification not in {"csv","folder"}: raise ValueError("Modo de identificação inválido.")
        if identification=="csv" and not skus: raise ValueError("Selecione a planilha no modo CSV ou utilize nome da pasta.")
        origin=settings.credentials(company,"wc").base_url
        identifier=uuid.uuid4().hex; folder=self.root/identifier
        folder.mkdir()
        try:
            found,embedded=extract_files(archive,folder,identification=identification,archive_name=archive_name,root=root,company=company,limit_mib=self.policy.max_file_mib,package_mib=self.policy.max_package_mib)
            if embedded:
                if complements is not None: raise ValueError("Use uma única planilha complementar: externa ou dentro do ZIP.")
                complements=parse_skus(*embedded,details=True,strict=True)
            if identification=="folder":
                if skus and any(s not in found for s in skus): raise ValueError("SKU da planilha não corresponde exatamente ao nome da pasta/ZIP.")
                skus=found
            complements=complements or {}
            if any(s not in skus for s in complements): raise ValueError("Planilha contém SKU sem pasta correspondente.")
            if any(r.get("empresa", "").strip() and r["empresa"].strip().lower()!=company for r in complements.values()):
                raise ValueError("Empresa da planilha diverge da loja selecionada.")
        except BaseException:
            # Rejected intake leaves no orphan extraction; the folder is new and owned by this call.
            shutil.rmtree(folder,ignore_errors=True)
            raise
        with self.db() as db:
            db.execute("INSERT INTO batches VALUES(?,?,?,?,?,?)",(identifier,str(org),str(owner),company,origin,time.time()))
            db.executemany("INSERT INTO jobs(batch,sku,state,plan) VALUES(?,?,'preview_queued',?)",
                [(identifier,s,json.dumps({"input":{"identification":identification,"complement":complements.get(s,{}),
                                                    **({"scope":scope} if scope else {})}})) for s in skus])
        return identifier

    def access(self, db, identifier, org, owner):
        row=db.execute("SELECT * FROM batches WHERE id=? AND org=? AND owner=?",(identifier,str(org),str(owner))).fetchone()
        if not row: raise LookupError("Lote não encontrado.")
        return row

    def audit(self, batch, sku, actor, event, detail=None, db=None):
        """Append-only audit trail; never stores paths, credentials or HTTP bodies."""
        row=(batch,sku,time.time(),str(actor),event,json.dumps(detail or {},ensure_ascii=False))
        if db is not None: db.execute("INSERT INTO audit VALUES(?,?,?,?,?,?)",row)
        else:
            with self.db() as conn: conn.execute("INSERT INTO audit VALUES(?,?,?,?,?,?)",row)

    def detail(self, identifier, org, owner):
        with self.db() as db:
            batch=dict(self.access(db,identifier,org,owner))
            items=[]
            for row in db.execute("SELECT * FROM jobs WHERE batch=?",(identifier,)).fetchall():
                plan=json.loads(row["plan"])
                dry=plan.get("dry_run")
                # No local paths, remote metadata or source URLs in frontend.
                items.append({"sku":row["sku"],"state":row["state"],"error":row["error"],
                    "files":[{"name":i["name"],"bytes":i["bytes"],"role":i.get("role","image"),"label":i.get("label","")}
                             for i in plan.get("images",[])+plan.get("downloads",[])]+
                            [{"name":i["name"],"bytes":i["bytes"],"role":"source","label":"fonte (não publicada)"} for i in plan.get("sources",[])]+
                            [{"name":i["name"],"bytes":i["bytes"],"role":i["role"],"label":"fora do escopo"} for i in plan.get("out_of_scope",[])],
                    "scope":plan.get("input",{}).get("scope"),"main_choice":plan.get("input",{}).get("main_choice"),
                    "existing_images":plan.get("existing_images"),"existing_documents":plan.get("existing_documents"),
                    "product_id":plan.get("product_id"),"dry_run":dry,
                    "blockers":(dry or {}).get("blockers",[]),"approval":plan.get("approval"),
                    "result":plan.get("result"),
                    "audit":[{"at":a["at"],"actor":a["actor"],"event":a["event"],"detail":json.loads(a["detail"])} for a in
                             db.execute("SELECT * FROM audit WHERE batch=? AND sku=? ORDER BY at",(identifier,row["sku"]))],
                    "preparation_status":"prepared" if plan.get("product_id") else "pending",
                    "warnings":plan.get("warnings",[])})
        for item in items:
            # Never expose old authorization jargon, and never mutate the persisted signed Dry Run.
            if item["dry_run"] and not item["dry_run"].get("authorization"):
                item["dry_run"] = {**item["dry_run"], "write_gate": "Dry Run aprovado. Confirme o envio." if approved(item) else "Revise e aprove o Dry Run antes de enviar."}
            item["write_gate"] = "Dry Run aprovado. Gestores e desenvolvedores podem confirmar o envio." if approved(item) else "Revise e aprove o Dry Run antes de enviar."
            item["needs_refresh"] = not (item["dry_run"] or {}).get("authorization")
            item["can_rollback"] = item["state"] in {"done", "uncertain", "error", "rollback_error"} and self.has_recovery(batch, item["sku"])
        ready=[i for i in items if i["state"]=="review" and approved(i) and not i["needs_refresh"]]
        return {"id":identifier,"company":batch["company"],"created":batch["created"],"items":items,"can_confirm":bool(ready)}

    def has_recovery(self, batch, sku):
        with self.db() as db:
            row=db.execute("SELECT plan FROM jobs WHERE batch=? AND sku=?", (batch["id"],sku)).fetchone()
        plan=json.loads(row["plan"]) if row else {}
        if not plan.get("dry_run"): return False
        entry=Journal(self,str(batch["org"])+"|"+batch["origin"]+"|"+sku).get(grants.update_key(plan["product_id"],plan["dry_run"]["dry_run_sha256"]))
        return bool(entry and entry.get("rollback"))

    def history(self,org,owner):
        with self.db() as db:
            return [dict(r) for r in db.execute("SELECT id,company,created FROM batches WHERE org=? AND owner=? ORDER BY created DESC LIMIT 100",(str(org),str(owner)))]

    def approve(self, identifier, org, owner, skus=None, replace_main=()):
        """Individual (skus) or batch approval of the exact Dry Run. Never queues or authorizes a write.

        Replacing an existing main image requires naming the SKU in replace_main.
        """
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            record=dict(self.access(db,identifier,org,owner))
            rows=[r for r in db.execute("SELECT * FROM jobs WHERE batch=? AND state='review'",(identifier,)).fetchall()
                  if skus is None or r["sku"] in skus]
            if skus is not None and len(rows)!=len(set(skus)): raise ValueError("SKU fora do lote ou sem Dry Run apto.")
            approved_skus=[]
            for row in rows:
                plan=json.loads(row["plan"])
                authorization.validate(record,row["sku"],plan,approval=False)
                main=plan["dry_run"]["images"]["main_image"]
                if plan["dry_run"].get("registration_scope"):
                    # Scoped plans already name the main file chosen for this SKU; approving the Dry Run approves it.
                    replace=main["action"]=="replace"
                else:
                    replace=row["sku"] in replace_main
                    if replace and not main["replacement_available"]: raise ValueError("Substituição de principal indisponível para "+row["sku"]+".")
                plan["approval"]={"by":str(owner),"at":time.time(),"dry_run_sha256":plan["dry_run"]["dry_run_sha256"],
                                  "replace_main":replace,"writes_authorized":False}
                db.execute("UPDATE jobs SET plan=? WHERE batch=? AND sku=? AND state='review'",(json.dumps(plan),identifier,row["sku"]))
                self.audit(identifier,row["sku"],owner,"approved",{"dry_run_sha256":plan["approval"]["dry_run_sha256"],"replace_main":replace},db)
                approved_skus.append(row["sku"])
            if not approved_skus: raise ValueError("Nenhum Dry Run pronto para aprovação.")
            return approved_skus

    def confirm(self, identifier, org, owner, settings, skus=None, *, confirmed=False, session_factory=None):
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            batch=dict(self.access(db,identifier,org,owner))
            if not confirmed: raise ValueError("Confirme explicitamente o envio dos arquivos.")
            if settings.organization_id != str(org) or settings.credentials(batch["company"],"wc").base_url!=batch["origin"]:
                raise ValueError("Organização ou destino alterado; importe um novo lote.")
            rows=[r for r in db.execute("SELECT * FROM jobs WHERE batch=? AND state='review'",(identifier,)).fetchall()
                  if approved(json.loads(r["plan"])) and (skus is None or r["sku"] in skus)]
            if skus is not None and len(rows)!=len(set(skus)):
                raise ValueError("SKU fora do lote ou sem Dry Run aprovado.")
            if not rows: raise ValueError("Nenhum item aprovado apto.")
            for row in rows:
                plan=json.loads(row["plan"])
                authorization.validate(batch,row["sku"],plan)
                authorization.check_identity(owner,org,batch["company"],plan["approval"]["by"],session_factory=session_factory)
                plan["execution"]=authorization.confirmation(batch,row["sku"],plan,owner,"publish")
                db.execute("UPDATE jobs SET state='queued',plan=?,error='' WHERE batch=? AND sku=?",
                           (json.dumps(plan),identifier,row["sku"]))
                self.audit(identifier,row["sku"],owner,"queued",{"authorization":authorization.VERSION,"dry_run_sha256":plan["dry_run"]["dry_run_sha256"]},db)

    def request_rollback(self, identifier, org, owner, skus, *, confirmed=False, session_factory=None):
        if not confirmed or not skus: raise ValueError("Selecione o SKU e confirme a restauração.")
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            batch=dict(self.access(db,identifier,org,owner))
            rows=[r for r in db.execute("SELECT * FROM jobs WHERE batch=?", (identifier,)).fetchall()
                  if r["sku"] in skus and r["state"] in {"done","uncertain","error","rollback_error"}]
            if len(rows)!=len(set(skus)): raise ValueError("SKU sem publicação apta à restauração.")
            for row in rows:
                plan=json.loads(row["plan"])
                authorization.validate(batch,row["sku"],plan,recovery=True)
                authorization.check_identity(owner,org,batch["company"],plan["approval"]["by"],session_factory=session_factory)
                key=grants.update_key(plan["product_id"],plan["dry_run"]["dry_run_sha256"])
                scope=batch["org"]+"|"+batch["origin"]+"|"+row["sku"]
                entry=db.execute("SELECT value FROM journal WHERE scope=? AND key=?", (scope,key)).fetchone()
                if not entry or not json.loads(entry[0]).get("rollback"):
                    raise ValueError("Não há publicação registrada para recuperar.")
                plan["rollback_request"]=authorization.confirmation(batch,row["sku"],plan,owner,"rollback")
                db.execute("UPDATE jobs SET state='rollback_queued',plan=?,error='' WHERE batch=? AND sku=?",
                           (json.dumps(plan),identifier,row["sku"]))
                self.audit(identifier,row["sku"],owner,"rollback_queued",{"authorization":authorization.VERSION},db)

    REPREVIEW_STATES=("error","awaiting_product","review","blocked","awaiting_main","nothing")

    def _requeue(self, db, batch, identifier, owner, rows, change=None, event="preview_requested", detail=None):
        """New Dry Run for these rows (approval dropped); change(input) may adjust scope/main first. Returns requeued SKUs."""
        done=[]
        for row in rows:
            plan=json.loads(row["plan"])
            if plan.get("dry_run"):
                key=grants.update_key(plan["product_id"],plan["dry_run"]["dry_run_sha256"])
                scope=batch["org"]+"|"+batch["origin"]+"|"+row["sku"]
                entry=db.execute("SELECT value FROM journal WHERE scope=? AND key=?", (scope,key)).fetchone()
                if entry and json.loads(entry[0]).get("state") in {"sending","done"}:
                    continue  # Do not lose recovery data for an uncertain association.
            inputs=dict(plan.get("input",{}))
            if change: change(inputs,plan)
            db.execute("UPDATE jobs SET state='preview_queued',error='',plan=? WHERE batch=? AND sku=?",
                       (json.dumps({"input":inputs}),identifier,row["sku"]))
            self.audit(identifier,row["sku"],owner,event,detail or {},db)
            done.append(row["sku"])
        return done

    def _rows(self, db, identifier, skus):
        marks=",".join("?"*len(self.REPREVIEW_STATES))
        return [r for r in db.execute(f"SELECT * FROM jobs WHERE batch=? AND state IN ({marks})",(identifier,*self.REPREVIEW_STATES)).fetchall()
                if skus is None or r["sku"] in skus]

    def reprocess(self, identifier, org, owner, skus=None):
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            batch=dict(self.access(db,identifier,org,owner))
            return self._requeue(db,batch,identifier,owner,self._rows(db,identifier,skus))

    def set_scope(self, identifier, org, owner, scope, skus=None):
        """Change what is registered for these items. Always a new Dry Run: a previous approval never carries over."""
        scope=normalize_scope(scope)
        def change(inputs,plan):
            inputs["scope"]=scope
            if not scope["main"]: inputs.pop("main_choice",None)
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            batch=dict(self.access(db,identifier,org,owner))
            rows=self._rows(db,identifier,skus)
            if skus is not None and len(rows)!=len(set(skus)): raise ValueError("SKU fora do lote ou em envio; escopo não alterado.")
            done=self._requeue(db,batch,identifier,owner,rows,change,"scope_changed",{"scope":scope})
            if not done: raise ValueError("Nenhum item apto a mudar de escopo.")
            return done

    def choose_main(self, identifier, org, owner, sku, sha256):
        """Explicit main image for one SKU, among the images of its own Dry Run. Triggers a new Dry Run."""
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            batch=dict(self.access(db,identifier,org,owner))
            rows=self._rows(db,identifier,[sku])
            if len(rows)!=1: raise ValueError("SKU fora do lote ou em envio; principal não alterada.")
            plan=json.loads(rows[0]["plan"])
            if not (plan.get("input",{}).get("scope") or {}).get("main"):
                raise ValueError("Marque Foto principal no escopo antes de escolher a imagem.")
            rule=(plan.get("dry_run") or {}).get("image_rule") or {}
            if rule.get("main"):
                raise ValueError(f"Principal definida pelo padrão {sku}_1.ext ({rule['main']['name']}); renomeie os arquivos para mudar.")
            option=next((o for o in (plan.get("dry_run") or {}).get("main_options",[]) if o["sha256"]==sha256),None)
            if not option: raise ValueError("Imagem não pertence a este SKU; escolha uma das imagens listadas no Dry Run.")
            choice={"sha256":option["sha256"],"name":option["name"]}
            done=self._requeue(db,batch,identifier,owner,rows,lambda inputs,_:inputs.update(main_choice=choice),"main_chosen",choice)
            if not done: raise ValueError("Envio anterior incerto para este SKU; reconciliar antes de mudar a principal.")
            return done

    def claim(self, batch=None, sku=None):
        """One item globally at a time. batch/sku restrict the claim to exactly that item (worker --batch/--sku)."""
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            db.execute("UPDATE jobs SET state='uncertain',error='Worker interrompido; reconciliar manualmente.',token='' WHERE state IN ('preview_running','running','rollback_running') AND lease<?",(time.time(),))
            if db.execute("SELECT 1 FROM jobs WHERE state IN ('preview_running','running','rollback_running')").fetchone(): return None
            where="".join([" AND j.batch=?" if batch else ""," AND j.sku=?" if sku else ""])
            row=db.execute("SELECT j.*,b.company,b.origin,b.org,b.owner FROM jobs j JOIN batches b ON b.id=j.batch WHERE state IN ('preview_queued','queued','rollback_queued')"+where+" ORDER BY b.created LIMIT 1",
                           tuple(v for v in (batch,sku) if v)).fetchone()
            if not row: return None
            job=dict(row);job["token"]=uuid.uuid4().hex
            db.execute("UPDATE jobs SET state=?,token=?,lease=? WHERE batch=? AND sku=?",
                ({"preview_queued":"preview_running","queued":"running","rollback_queued":"rollback_running"}[job["state"]],job["token"],time.time()+3600,job["batch"],job["sku"]))
            return job

    def finish(self, job, state, plan=None, error=""):
        with self.db() as db:
            db.execute("UPDATE jobs SET state=?,plan=?,error=?,token='' WHERE batch=? AND sku=? AND token=?",
                (state,json.dumps(plan or json.loads(job["plan"])),error,job["batch"],job["sku"],job["token"]))


AWAITING_PRODUCT="Aguardando cadastro por CSV: SKU não existe na loja; criação automática proibida. Retome após o cadastro."


def approved(item):
    approval,dry=item.get("approval"),item.get("dry_run") or {}
    return bool(approval) and bool(dry.get("dry_run_sha256")) and approval.get("dry_run_sha256")==dry.get("dry_run_sha256")


class Journal:
    def __init__(self,jobs,scope): self.jobs,self.scope=jobs,scope
    def get(self,key):
        with self.jobs.db() as db:
            row=db.execute("SELECT value FROM journal WHERE scope=? AND key=?",(self.scope,key)).fetchone()
            return json.loads(row[0]) if row else None
    def put(self,key,value):
        with self.jobs.db() as db:
            db.execute("INSERT OR REPLACE INTO journal VALUES(?,?,?)",(self.scope,key,json.dumps(value)))


def unchanged_files(jobs, job, plan):
    """Before any upload: the SKU folder still holds exactly the files (names, sizes, bytes) of the approved Dry Run."""
    changed=ValueError("Arquivos do SKU renomeados, removidos ou adicionados após o Dry Run aprovado; refaça a prévia. Nada enviado.")
    folder=find_product_folder(jobs.root/job["batch"],job["sku"])
    if not folder or folder_inventory(folder)!=plan["dry_run"]["files_inventory_sha256"]: raise changed
    for item in plan["images"]+plan["downloads"]:
        try:
            if file_info(item["path"],jobs.policy.max_file_mib)["sha256"]!=item["sha256"]: raise changed
        except ValueError:
            raise changed from None


def process_one(jobs, settings, transport=None, *, batch=None, sku=None, session_factory=None):
    job=jobs.claim(batch,sku)
    if not job: return False
    woo=None;publisher=None
    try:
        if settings.organization_id!=job["org"]:
            raise ValueError("Organização não corresponde à configuração.")
        if settings.credentials(job["company"],"wc").base_url!=job["origin"]:
            raise ValueError("Destino mudou; importe novo lote.")
        if job["state"]=="preview_queued":
            inputs=json.loads(job["plan"]).get("input",{})
            folder=find_product_folder(jobs.root/job["batch"],job["sku"])
            if not folder: raise ValueError("Pasta exata do SKU não encontrada no ZIP.")
            if folder.name!=job['sku']: raise ValueError("Nome da pasta não corresponde exatamente ao SKU.")
            files=prepare_files(folder,job["sku"],company=job["company"],numbered=inputs.get("scope") is not None)
            selected=([files["main_image"]] if files["main_image"] else [])+files["gallery"]
            images=[{"path":str(p),"role":"main_image" if n==0 else "gallery",**file_info(p,jobs.policy.max_file_mib)}
                    for n,p in enumerate(selected)]
            if len(images)>20: raise ValueError("Limite de 20 imagens por produto.")
            # Formats not homologated for upload (e.g. .step/.dwg/.rar) are reported and ignored, never sent.
            unsupported=[d for d in files["downloads"] if Path(d["path"]).suffix.lower() not in MIMES]
            documents=[{**d,**file_info(d["path"],jobs.policy.max_file_mib)} for d in files["downloads"] if d not in unsupported]
            if not images and not documents: raise ValueError("Nenhum arquivo reconhecido.")
            # The scope and the chosen main survive every new Dry Run of this item.
            kept={k:inputs[k] for k in ("identification","scope","main_choice") if inputs.get(k) is not None}
            scope=inputs.get("scope")
            out_of_scope,scope_ignored,options,rule=[],[],[],None
            if scope is not None:
                scope=normalize_scope(scope)
                candidates=images
                if files["manifest"]:
                    # A validated PRM manifest keeps its declared order; the main is chosen in the Dry Run.
                    options=main_options(images,files["main_from_pattern"])
                    images,scope_ignored=scoped_images(candidates,scope,inputs.get("main_choice"),files["main_from_pattern"])
                else:
                    # Raw folders: <SKU>_1.ext is the main, _2.. the gallery. Without _1, only a manual choice.
                    rule=classify_by_index(candidates,job["sku"])
                    options=[] if rule["main"] else main_options(candidates,False)
                    images,scope_ignored=rule_images(rule,scope,inputs.get("main_choice"),job["sku"])
                scope_ignored+=[{"name":i["name"],"label":"Formato não enviado",
                                 "reason":"SVG/GIF não é enviado pelo cadastro via API; use a versão PNG/JPG/WEBP (Resolução de Imagens)."}
                                for i in files.get("unsupported_images",[])]
                planned={i["path"] for i in images}
                out_of_scope=[{"name":i["name"],"bytes":i["bytes"],"role":"image"} for i in candidates if i["path"] not in planned]
                if not scope["documents"]:
                    out_of_scope+=[{"name":d["name"],"bytes":d["bytes"],"role":d["role"]} for d in documents]
                    documents=[]
            woo=WooClient(credentials_for(job["company"],settings),transport=transport)
            product=woo.find_by_sku(job["sku"])
            if not product:
                # Never created here: the product is registered by the CSV flow, then the item is resumed.
                jobs.finish(job,"awaiting_product",{"input":kept,"sources":files.get("sources",[]),
                    "warnings":files["warnings"]},AWAITING_PRODUCT)
                jobs.audit(job["batch"],job["sku"],"worker","awaiting_product")
                return True
            product=woo.product(product["id"])
            if product.get('sku')!=job['sku']: raise ValueError("SKU mudou durante a consulta; refaça a prévia.")
            # Private immutable backup for this preview attempt. No remote write is performed.
            backup=jobs.root/'_snapshots'/job['batch']
            backup.mkdir(parents=True,exist_ok=True)
            snapshot=json.dumps({"company":job['company'],"origin":job['origin'],"sku":job['sku'],"product":product},ensure_ascii=False)
            with (backup/(job['token']+'.json')).open('x',encoding='utf-8') as handle: handle.write(snapshot)
            dry={"dry_run":True,"writes_allowed":False,"product_id":product["id"],"sku":job["sku"],
                 **(file_plan(product,images,documents) if scope is None else
                    scoped_file_plan(product,images,documents,scope,inputs.get("main_choice"),options,rule)),
                 "backup_sha256":hashlib.sha256(snapshot.encode('utf-8')).hexdigest(),
                 "ignored":[{"name":Path(d["path"]).name,"label":d["label"],"reason":"Formato não suportado para upload."} for d in unsupported]+
                           [{"name":d["name"],"label":"Datasheet do fornecedor","reason":"Datasheet original do fornecedor não é publicado; somente o datasheet gerado pelo PRM."}
                            for d in files.get("excluded",[])]+scope_ignored}
            if scope is not None:
                # Signed with the Dry Run: renaming, adding or removing any file of the SKU invalidates the execution.
                dry["files_inventory_sha256"]=folder_inventory(folder)
            # Store-side limit (WordPress/PHP), read-only; checked before approval so nothing is sent in vain.
            limit_warnings=store_limit_check(dry,woo.upload_limit())
            dry["write_gate"]="Revise e aprove o Dry Run; gestores e desenvolvedores podem confirmar o envio."
            current_documents=interpret_documents(product)
            plan={"input":kept if scope is not None else {"identification":inputs.get("identification")},
                  "product_id":product["id"],"snapshot_hash":snapshot_hash(product),"protected_hash":_protected_hash(product),
                  "images":images,"downloads":documents,"sources":files.get("sources",[]),"out_of_scope":out_of_scope,
                  "existing_images":len(product.get("images",[])),
                  "existing_documents":len(current_documents["documents"]),
                  "dry_run":dry,
                  "warnings":files["warnings"]+current_documents["warnings"]+limit_warnings}
            authorization.seal(job,job["sku"],plan)
            if dry["blockers"]:
                # Only the missing main choice left: the user picks it per SKU (not an error).
                state="awaiting_main" if dry.get("main_choice_required") and len(dry["blockers"])==1 else "blocked"
            else:
                state="nothing" if dry.get("nothing_to_do") else "review"
            jobs.finish(job,state,plan,"; ".join(dry["blockers"]) if dry["blockers"] else dry.get("nothing_message") or "")
            jobs.audit(job["batch"],job["sku"],"worker","preview_"+state,{"dry_run_sha256":dry["dry_run_sha256"]})
        else:
            plan=json.loads(job["plan"])
            journal=Journal(jobs,job["org"]+"|"+job["origin"]+"|"+job["sku"])
            action="rollback" if job["state"]=="rollback_queued" else "publish"
            if action=="publish" and stale_supplier_datasheet(plan["dry_run"]):
                raise ValueError("Dry Run anterior à regra do datasheet do fornecedor; refaça a prévia e aprove novamente. Nada enviado.")
            if action=="publish" and plan["dry_run"].get("files_inventory_sha256"):
                unchanged_files(jobs,job,plan)
            check=lambda: authorization.recheck(jobs,job,plan,action,session_factory)
            publisher=StagingPublisher(job["company"],job["sku"],settings,jobs.policy,journal,transport=transport,authorization_check=check)
            actor=(plan.get("execution" if action=="publish" else "rollback_request") or {}).get("by","worker")
            audit=lambda event,detail=None:jobs.audit(job["batch"],job["sku"],actor,event,detail)
            result=publisher.rollback(plan,audit=audit) if action=="rollback" else publisher.execute(plan,audit=audit)
            jobs.finish(job,"rolled_back" if action=="rollback" else "done",{**plan,("rollback_result" if action=="rollback" else "result"):result})
            audit("rollback_done" if action=="rollback" else "done",{"verified":result.get("verified")})
    except Exception as exc:
        # HTTP bodies and local filesystem paths must not leak to reports.
        error=str(exc) if isinstance(exc,(ValueError,authorization.access.PermissionDenied)) and len(str(exc))<300 else "Falha "+type(exc).__name__+"; confira configuração e arquivos."
        state="rollback_error" if job["state"]=="rollback_queued" else "error"
        if job["state"]=="queued":
            p=json.loads(job["plan"])
            if p.get("dry_run"):
                entry=Journal(jobs,job["org"]+"|"+job["origin"]+"|"+job["sku"]).get(grants.update_key(p["product_id"],p["dry_run"]["dry_run_sha256"]))
                if entry and entry.get("state")=="sending": state="uncertain"
        jobs.finish(job,state,error=error)
        jobs.audit(job["batch"],job["sku"],"worker","error",{"error":error})
    finally:
        if woo: woo.close()
        if publisher: publisher.close()
    return True


def rollback_one(jobs, settings, batch, sku, *, actor, confirmed=False, transport=None, session_factory=None):
    """Compatibility helper: same authenticated confirmation/queue as UI, no terminal grants."""
    if not confirmed: raise ValueError("Confirme a restauração pela interface autenticada.")
    jobs.request_rollback(batch,settings.organization_id,actor,[sku],confirmed=True,session_factory=session_factory)
    process_one(jobs,settings,transport,batch=batch,sku=sku,session_factory=session_factory)
    item=jobs.detail(batch,settings.organization_id,actor)["items"]
    row=next(i for i in item if i["sku"]==sku)
    if row["state"]!="rolled_back": raise ValueError(row["error"] or "Restauração pendente na fila.")
    with jobs.db() as db:
        plan=json.loads(db.execute("SELECT plan FROM jobs WHERE batch=? AND sku=?",(batch,sku)).fetchone()[0])
    return plan["rollback_result"]
