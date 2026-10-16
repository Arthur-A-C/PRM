"""Catálogos refinados versionados, independentes do banco e sem rede/IA.

Nesta entrega famílias são propostas e cobertura técnica é não avaliada.
Nenhum resultado legado ECO é usado como característica de normalização.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import uuid
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from pathlib import Path
from time import perf_counter

from openpyxl import load_workbook

from app.services.catalog_semantics import normalize_cpu
from app.services.operating_system import os_families, combined_os, VERSION as OS_VERSION
from app.services.catalog_attribute_analysis import AttributeAnalysis
from app.services.attribute_dictionary import apply_alias, attribute_name, load_dictionary
from app.services.catalog_treatment import plain
from app.services.dimensional_csv import HEADERS, sha256_file, valid_gtin

ROOT = Path(__file__).resolve().parents[2]
POLICY = ROOT / "config/catalogs/policy.json"
STAGING = ROOT / "data/analysis/staging/refined"
ARCHIVE = ROOT / "data/imports/refined"
NAMESPACE = uuid.UUID("9b91e2e4-cdfe-4b2f-a84f-f8c88bf78c58")
ORDER = ["portwell", "choice", "eco", "dimensional"]


def dumps(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False, default=str)


def fingerprint(value):
    return hashlib.sha256(dumps(value).encode("utf-8")).hexdigest()


class _Text(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.skip += 1
        elif tag in {"br", "p", "div", "li", "tr"}:
            self.parts.append(" ")

    def handle_endtag(self, tag):
        if tag in {"script", "style"} and self.skip:
            self.skip -= 1

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def clean_text(value):
    parser = _Text()
    parser.feed(str(value or ""))
    return " ".join(" ".join(parser.parts).split())


def load_policy(path=POLICY):
    policy = json.loads(Path(path).read_text(encoding="utf-8"))
    if policy.get("schema_version") != 1 or set(policy.get("sources", {})) != set(ORDER):
        raise ValueError("Política de origens incompatível")
    if policy["catalog_policy"].get("absent_in_new_file_means_discontinued") is not False:
        raise ValueError("Descontinuação automática não homologada")
    if policy["matching"].get("enabled") or policy["ai"].get("enabled") or policy["documents"].get("automatic_download_enabled"):
        raise ValueError("Esta entrega não habilita matching, IA paga ou downloads")
    ids = [f["id"] for f in policy["families"]]
    if len(set(ids)) != len(ids):
        raise ValueError("Famílias duplicadas")
    for family in policy["families"]:
        re.compile(family["pattern"])
        if family["proximity"] not in {"core", "adjacent", "distant", "undecided"}:
            raise ValueError("Proximidade inválida")
    policy["attribute_dictionary"] = load_dictionary()
    policy["operating_system_policy"] = OS_VERSION
    return policy


def load_inputs(manifest):
    doc = json.loads(Path(manifest).read_text(encoding="utf-8"))
    entries = doc.get("inputs", [])
    if doc.get("schema_version") != 1 or sorted(e.get("source_id", "") for e in entries) != sorted(ORDER):
        raise ValueError("Manifesto deve conter as quatro origens exatamente uma vez")
    result = []
    for entry in entries:
        path = Path(entry["path"])
        if not path.is_absolute():
            path = ROOT / path
        result.append({"source_id": entry["source_id"], "path": path.resolve(strict=True)})
    return sorted(result, key=lambda e: ORDER.index(e["source_id"]))


def headers_for(path, source):
    if source == "eco":
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            return [str(c.value or "") for c in next(book["Base_ECO"].iter_rows())]
        finally:
            book.close()
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return next(csv.reader(stream, strict=True))


def iter_records(path, source):
    """Todos os registros são contados; fórmulas nunca são executadas."""
    headers = headers_for(path, source)
    required = {"Codigo_ECO", "Nome_Original"} if source == "eco" else set(HEADERS) if source == "dimensional" else {"ID", "Tipo", "SKU", "Nome", "Publicado", "Visibilidade no catálogo"}
    if required - set(headers) or len(headers) != len(set(headers)) or "" in headers:
        raise ValueError(f"{source}: cabeçalhos incompatíveis ou duplicados")
    if source == "eco":
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            for number, cells in enumerate(book["Base_ECO"].iter_rows(min_row=2), 1):
                yield {"position": number, "first_line": number + 1, "last_line": number + 1,
                       "raw": {h: c.value for h, c in zip(headers, cells)},
                       "cell_types": {h: c.data_type for h, c in zip(headers, cells) if c.value is not None},
                       "formula_fields": [h for h, c in zip(headers, cells) if c.data_type == "f"],
                       "issues": [], "error": None}
        finally:
            book.close()
        return
    csv.field_size_limit(4 * 1024 * 1024)
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream, strict=True)
        next(reader)
        previous = reader.line_num
        for position, values in enumerate(reader, 1):
            width, issues, error = len(values), [], None
            if width > len(headers) or (source == "dimensional" and width != len(headers)):
                error = "record_width_mismatch"
            elif width < len(headers):
                end = headers.index("Categorias") if "Categorias" in headers else len(headers)
                valid = (width > end and values[headers.index("ID")].isdigit()
                         and values[headers.index("Tipo")] in {"simple", "variable", "variation", "grouped", "external"}
                         and values[headers.index("Publicado")] in {"1", "0", "-1"}
                         and values[headers.index("Visibilidade no catálogo")] in {"visible", "hidden", "catalog", "search"}
                         and bool(values[headers.index("Nome")].strip()))
                if valid:
                    issues.append("trailing_fields_padded")
                else:
                    error = "short_record_not_validated"
            raw = dict(zip(headers, values + [None] * max(0, len(headers) - width)))
            if source in {"portwell", "choice"} and not error:
                if (not str(raw.get("ID", "")).isdigit()
                        or raw.get("Tipo") not in {"simple", "variable", "variation", "grouped", "external"}
                        or raw.get("Publicado") not in {"1", "0", "-1"}
                        or raw.get("Visibilidade no catálogo") not in {"visible", "hidden", "catalog", "search"}):
                    error = "structural_alignment_invalid"
            if width > len(headers):
                raw["_extra_values"] = values[len(headers):]
            yield {"position": position, "first_line": previous + 1, "last_line": reader.line_num,
                   "raw": raw, "issues": issues, "error": error, "field_count": width}
            previous = reader.line_num


def text_value(raw, field, formulas=()):
    value = raw.get(field)
    return "" if value is None or field in formulas else str(value).strip()


def parse_price(raw_value, locale):
    raw = str(raw_value or "").strip()
    if not raw:
        return None, "unknown"
    number = raw.replace(".", "").replace(",", ".") if locale == "pt_BR" else raw
    if not re.fullmatch(r"\d+(?:\.\d+)?", number):
        return None, "invalid"
    try:
        amount = Decimal(number)
    except InvalidOperation:
        return None, "invalid"
    return (str(amount), "known") if amount > 0 else (None, "zero_not_free")


def split_pairs(value, separator=":"):
    pairs, residuals = [], []
    for segment in str(value or "").split("|"):
        segment = segment.strip()
        if not segment:
            continue
        if separator not in segment:
            residuals.append(segment)
            continue
        key, val = segment.split(separator, 1)
        if key.strip() and val.strip():
            pairs.append((key.strip(), val.strip()))
        else:
            residuals.append(segment)
    return pairs, residuals


def classify(title, category, policy):
    for text in [plain(title), plain(category)]:
        for family in policy["families"]:
            if re.search(family["pattern"], text):
                return {"family_id": family["id"], "family_label": family["label"],
                        "item_form": family.get("form", "unknown"), "proximity": family["proximity"],
                        "verification_state": "unreviewed", "method": "deterministic_seed"}
    return {"family_id": "unknown", "family_label": "Família pendente", "item_form": "unknown",
            "proximity": "undecided", "verification_state": "unreviewed", "method": "unresolved"}


def normalize_attribute(name, raw_value, family, title=""):
    label = attribute_name(name)
    value = str(raw_value).strip().replace("\\,", ",")
    if value.startswith("'-") and re.search(r"\d", value):
        value = value[1:]
    result = {"attribute_code": "unmapped." + label.replace(" ", "_"), "normalized_value": value,
              "value_type": "text", "unit": None, "qualifiers": {}, "configuration_scope": "unspecified",
              "knowledge_state": "unknown" if value in {"", "-", "'-"} else "known",
              "verification_state": "unreviewed", "mapped": False, "issues": []}
    if label in {"protocolos de comunicacao", "protocolo", "protocolo de comunicacao"}:
        result.update(attribute_code="communication.protocol", mapped=True)
        if "rj45" in plain(value) and not re.search(r"modbus|profinet|ethercat|canopen", value, re.I):
            result.update(normalized_value=None, knowledge_state="unknown", qualifiers={"connector_observed": "RJ45"})
            result["issues"].append("connector_in_protocol_field")
        elif plain(value) in {"modbus tcp", "modbus tcp ip"}:
            result.update(normalized_value="modbus_tcp", value_type="enum")
    elif re.fullmatch(r"LGA\s*1700", value, re.I):
        result.update(attribute_code="compute.socket", normalized_value="LGA1700", value_type="enum", mapped=True)
    elif "processador" in label and normalize_cpu(value):
        result.update(attribute_code="compute.cpu.model", normalized_value=normalize_cpu(value), value_type="identifier", mapped=True)
    elif "rj45" in plain(value) and re.search(r"10/100\s*Mbps.*Gigabit", value, re.I):
        result.update(attribute_code="network.ports", normalized_value=None, knowledge_state="conflicting", mapped=True)
        result["issues"].append("ethernet_speed_conflict")
    elif "usb" in plain(value) and len(re.findall(r"\d+\s*x\s*USB", value, re.I)) >= 4:
        result.update(attribute_code="ports.usb", normalized_value=None, knowledge_state="conflicting", mapped=True)
        result["issues"].append("ambiguous_usb_quantities")
    elif label == "memoria":
        if family in {"box_pc", "panel_pc"}:
            result.update(attribute_code="compute.ram", mapped=True)
        elif family == "plc":
            result.update(attribute_code="control.memory", mapped=True)
        else:
            result.update(knowledge_state="unparseable")
    elif label in {"alcance", "distancia maxima", "distancia"}:
        if "μg/m³" in value or "µg/m³" in value:
            result.update(attribute_code="measurement.concentration_range", unit="µg/m³", mapped=True)
        elif re.search(r"\b(km|mm|m)\b", value):
            result.update(attribute_code="measurement.distance_range", unit=re.search(r"\b(km|mm|m)\b", value).group(1), mapped=True)
        else:
            result.update(knowledge_state="unparseable")
    elif "bitola" in label or "secao transversal" in label:
        result.update(attribute_code="conductor.cross_section", mapped=True)
        if re.search(r"\bmm\b(?![²2])", value) and not re.search(r"mm[²2]", value):
            result.update(knowledge_state="conflicting", normalized_value=None)
            result["issues"].append("area_unit_mismatch")
    elif "potencia" in label and re.search(r"\d+,\s+\d+\s*W", value):
        result.update(knowledge_state="conflicting")
        result["issues"].append("ambiguous_decimal_power")
        power = re.search(r"\d+,\d+\s*W", title, re.I)
        if power:
            result["resolution_proposal"] = {"value": power.group(0), "field": "title", "verification_state": "unreviewed"}
    if re.search(r"suporta|suportado", value, re.I):
        result["configuration_scope"] = "supported"
    elif re.search(r"opcional", value, re.I):
        result["configuration_scope"] = "optional"
    elif re.search(r"instalado", value, re.I):
        result["configuration_scope"] = "installed"
    return apply_alias(result, name, family)


def normalize_record(record, source, policy, batch_id):
    raw, formulas = record["raw"], record.get("formula_fields", [])
    get = lambda field: text_value(raw, field, formulas)
    cfg, issues = policy["sources"][source], list(record["issues"])
    key = get(cfg["key"])
    if not key:
        issues.append("missing_source_key")
        key = "provisional:" + fingerprint(raw)
    item_id = str(uuid.uuid5(NAMESPACE, source + ":" + key))
    title = get("Nome_Original" if source == "eco" else "nome" if source == "dimensional" else "Nome")
    sku = get("sku" if source == "dimensional" else "SKU")
    if cfg["kind"] == "own" and not sku:
        issues.append("empty_sku_preserved")
    if sku.casefold() == "codmaterialaqui":
        issues.append("placeholder_sku")
    if not title:
        issues.append("missing_title")
    category = get("Categoria_Site" if source == "eco" else "categoria" if source == "dimensional" else "Categorias")
    classification = classify(title, category, policy)
    brand = get("Marca_Padronizada" if source == "eco" else "marca" if source == "dimensional" else "Marcas")
    public = cfg["kind"] != "own" or (get("Publicado") == "1" and get("Visibilidade no catálogo") == "visible")
    pairs, residuals = [], []
    if cfg["kind"] == "own":
        for field in raw:
            match = re.fullmatch(r"Nome do atributo (\d+)", field)
            if match and get(field):
                value_field = "Valores do atributo " + match.group(1)
                if get(value_field):
                    pairs.append((get(field), get(value_field), value_field))
    elif source == "eco":
        pairs.extend((field, get(field), field) for field in raw if field.startswith("ATR_") and get(field))
        for field in ["Atributos_PRM", "Outros_Atributos"]:
            extra, residues = split_pairs(get(field), "=")
            pairs.extend((name, value, field) for name, value in extra)
            residuals.extend(residues)
    else:
        for field in ["especificacoes", "descricao_curta"]:
            extra, residues = split_pairs(get(field))
            pairs.extend((name, value, field) for name, value in extra)
            residuals.extend(residues)
    attributes, seen = [], {}
    for name, value, field in pairs:
        attr = normalize_attribute(name, value, classification["family_id"], title)
        evidence = {"source_id": source, "batch_id": batch_id, "record_key": key, "field": field,
                    "quote": value, "source_layer": "intermediate_eco_table" if source == "eco" else "source_export"}
        attr.update(item_id=item_id, raw_name=name, raw_value=value, evidence=[evidence],
                    method="deterministic", rule_version=policy["rule_version"])
        identity = (attr["attribute_code"], dumps(attr["normalized_value"]), attr["knowledge_state"],
                    attr["configuration_scope"], attr["unit"])
        if source == "eco" and identity in seen:
            seen[identity]["evidence"].append(evidence)
            seen[identity]["correlated_representations"] = True
            continue
        seen[identity] = attr
        attributes.append(attr)
        issues.extend(attr["issues"])
    by_code = {}
    for attr in attributes:
        by_code.setdefault((attr["attribute_code"], attr["configuration_scope"], attr["unit"]), []).append(attr)
    for observations in by_code.values():
        if observations[0]['attribute_code'] == 'compute.os' and combined_os(observations) is not None:
            for attr in observations:
                attr['qualifiers']['os_families'] = sorted(os_families(attr['raw_value']))
                attr['qualifiers']['comparison_policy'] = OS_VERSION
            continue
        if len({dumps(a["normalized_value"]) for a in observations if a["knowledge_state"] == "known"}) > 1:
            for attr in observations:
                attr["knowledge_state"] = "conflicting"
            issues.append("divergent_attribute_observations")
    if residuals:
        issues.append("attribute_parsing_residuals")
    price_fields = [("Preco_pix", "pix"), ("Preco_cartao", "card")] if source == "eco" else [("preco", "current"), ("preco_de", "reference")] if source == "dimensional" else [("Preço normal" if "Preço normal" in raw else "Preço", "regular"), ("Preço promocional", "sale")]
    hidden = cfg["kind"] == "own" and get("Metadado: _hide_product_price").casefold() in {"1", "yes", "true"}
    prices = []
    for field, kind in price_fields:
        amount, state = parse_price(get(field), cfg["locale"])
        prices.append({"raw_value": raw.get(field), "amount": amount, "state": state, "kind": kind,
                       "currency": "BRL", "currency_assumed": True, "visibility": "internal_only" if hidden else "source_declared",
                       "commercial_comparability": "not_verified", "source_observed_at": None, "field": field})
        if state in {"invalid", "zero_not_free"}:
            issues.append("price_" + state)
    availability = get("Estoque" if source == "eco" else "disponibilidade" if source == "dimensional" else "Em estoque?")
    stock = get("estoque" if source == "dimensional" else "Estoque") if source != "eco" else None
    if source == "dimensional" and stock == "99999":
        issues.append("stock_sentinel")
    gtin = get("ean" if source == "dimensional" else "GTIN, UPC, EAN, or ISBN" if "GTIN, UPC, EAN, or ISBN" in raw else "GTIN, UPC, EAN ou ISBN")
    if gtin and not valid_gtin(gtin):
        issues.append("invalid_gtin")
    model = get("Modelo_Referencia" if source == "eco" else "referencia" if source == "dimensional" else "")
    if model.count("(") != model.count(")"):
        issues.append("model_unbalanced_parentheses")
    if formulas:
        issues.append("excel_formula_not_evaluated")
    if classification["family_id"] == "unknown":
        issues.append("family_unmapped")
    item = {"item_id": item_id, "source_id": source, "record_key": key, "sku": sku or None,
            "url": get("Link" if source == "eco" else "url" if source == "dimensional" else "URL externa"),
            "description_original": get("descricao_completa" if source == "dimensional" else "Descrição"),
            "short_description_original": get("descricao_curta" if source == "dimensional" else "Descrição curta"),
            "source_page_id": get("product_id"), "title_original": title, "title_display": clean_text(title),
            "title_search": plain(clean_text(title)), "brand_claimed": brand or None,
            "brand_search": policy["brand_aliases"].get(plain(brand), plain(brand)) or None,
            "manufacturer_confirmed": None, "model_claimed": model or None, "gtin_raw": gtin or None,
            "gtin_valid": bool(gtin and valid_gtin(gtin)), "category_original": category,
            "classification": classification, "public_catalog": public, "catalog_completeness": cfg["completeness"],
            "commercial": {"prices": prices, "availability_raw": availability, "stock_raw": stock, "quantity": None,
                           "quantity_reliability": "unverified_source_semantics" if source == "dimensional" else "not_verified",
                           "freshness": "freshness_unknown"},
            "quality_issues": sorted(set(issues)), "attribute_count": len(attributes), "parsing_residuals": residuals,
            "batch_id": batch_id, "technical_match_status": "not_assessed", "substitution_approval": "not_reviewed"}
    item["normalized_hash"] = fingerprint({"item": item, "attributes": attributes})
    return item, attributes


def profile_inputs(entries, policy):
    reports = {}
    for entry in entries:
        source, path = entry["source_id"], entry["path"]
        headers, counters, widths, keys, attrs = headers_for(path, source), Counter(), Counter(), Counter(), Counter()
        for record in iter_records(path, source):
            counters["received"] += 1
            widths[str(record.get("field_count", len(headers)))] += 1
            if record["error"]:
                counters["quarantine"] += 1
                continue
            counters["accepted"] += 1
            counters.update(record["issues"])
            raw = record["raw"]
            keys[text_value(raw, policy["sources"][source]["key"], record.get("formula_fields", []))] += 1
            if source in {"choice", "portwell"}:
                counters["empty_sku"] += not bool(text_value(raw, "SKU"))
                counters["public_catalog"] += text_value(raw, "Publicado") == "1" and text_value(raw, "Visibilidade no catálogo") == "visible"
                has_attrs = False
                for field in raw:
                    match = re.fullmatch(r"Nome do atributo (\d+)", field)
                    if match and text_value(raw, field) and text_value(raw, "Valores do atributo " + match.group(1)):
                        attrs[text_value(raw, field)] += 1
                        has_attrs = True
                counters["without_attributes"] += not has_attrs
            counters["formula_records"] += bool(record.get("formula_fields"))
        reports[source] = {"file": path.name, "sha256": sha256_file(path), "size_bytes": path.stat().st_size,
                           "headers": headers, "schema_fingerprint": fingerprint(headers), "field_widths": dict(widths),
                           "counts": dict(counters), "distinct_keys": len(keys), "missing_keys": keys[""],
                           "duplicate_keys": {k: n for k, n in keys.items() if n > 1}, "attribute_labels": dict(attrs)}
    return {"schema_version": 1, "mode": "read_only_staging", "sources": reports,
            "received": sum(r["counts"]["received"] for r in reports.values())}


def spreadsheet_safe(value):
    raw = str(value or "")
    return "'" + raw if raw.lstrip().startswith(("=", "+", "-", "@")) else raw


def atomic_json(path, document):
    temp = path.with_name(path.name + "." + uuid.uuid4().hex + ".tmp")
    temp.write_text(dumps(document) + "\n", encoding="utf-8")
    temp.replace(path)


def run_pipeline(entries, policy, *, output_root=STAGING, archive_root=ARCHIVE):
    started = perf_counter()
    attribute_analysis = AttributeAnalysis()
    inputs = [{"source_id": e["source_id"], "file": e["path"].name, "sha256": sha256_file(e["path"])} for e in entries]
    run_id = fingerprint({"inputs": inputs, "policy": policy})[:24]
    output_root, archive_root = Path(output_root), Path(archive_root)
    output_root.mkdir(parents=True, exist_ok=True)
    target = output_root / run_id
    if (target / "run_manifest.json").is_file():
        manifest = verify_run(target)
        atomic_json(output_root / "latest.json", {"run_id": run_id})
        return {**manifest, "idempotent": True}
    work = output_root / (run_id + ".building")
    work.mkdir(exist_ok=True)
    atomic_json(work / "data_profile.json", profile_inputs(entries, policy))
    atomic_json(work / "policy.json", policy)
    outputs = {name: (work / name).open("w", encoding="utf-8", newline="") for name in ["raw_records.jsonl", "own_items.jsonl", "competitor_items.jsonl", "attribute_observations.jsonl", "quality_issues.csv", "competitor_coverage.csv", "legacy_eco_audit.csv", "quarantine.jsonl"]}
    issue_writer = csv.writer(outputs["quality_issues.csv"])
    coverage_writer = csv.writer(outputs["competitor_coverage.csv"])
    legacy_writer = csv.writer(outputs["legacy_eco_audit.csv"])
    issue_writer.writerow(["source_id", "record_key", "position", "code", "severity", "review_required"])
    coverage_writer.writerow(["run_id", "competitor_item_id", "source_id", "record_key", "family", "scope", "catalog_completeness", "coverage_status", "search_completed", "reason"])
    legacy_writer.writerow(["record_key", "legacy_result", "legacy_choice_code", "present_in_current_choice_skus", "new_status"])
    choice_skus, counts, quality, families, unmapped = set(), {}, {}, {}, Counter()
    for entry in entries:
        if entry["source_id"] == "choice":
            choice_skus = {text_value(r["raw"], "SKU") for r in iter_records(entry["path"], "choice") if not r["error"]}
    try:
        for entry, info in zip(entries, inputs):
            source, path = entry["source_id"], entry["path"]
            dest = archive_root / info["sha256"] / path.name
            dest.parent.mkdir(parents=True, exist_ok=True)
            if not dest.exists():
                copying = dest.with_suffix(dest.suffix + ".copying")
                shutil.copyfile(path, copying)
                if sha256_file(copying) != info["sha256"]:
                    raise ValueError("Entrada mudou durante cópia")
                copying.replace(dest)
            if sha256_file(dest) != info["sha256"]:
                raise ValueError("Hash do arquivo arquivado divergente")
            info["archived_path"] = str(dest.resolve())
            counters, issues, seen = Counter(), Counter(), set()
            batch_id = fingerprint({"source": source, "hash": info["sha256"]})[:24]
            for record in iter_records(dest, source):
                counters["received"] += 1
                outputs["raw_records.jsonl"].write(dumps({**record, "source_id": source, "batch_id": batch_id, "raw_hash": fingerprint(record["raw"])}) + "\n")
                key = text_value(record["raw"], policy["sources"][source]["key"])
                if key and key in seen:
                    record["error"] = "duplicate_source_key"
                if key:
                    seen.add(key)
                if record["error"]:
                    counters["quarantine"] += 1
                    outputs["quarantine.jsonl"].write(dumps({**record, "source_id": source}) + "\n")
                    issue_writer.writerow([source, spreadsheet_safe(key), record["position"], record["error"], "record_error", True])
                    continue
                item, attributes = normalize_record(record, source, policy, batch_id)
                attribute_analysis.observe(item, attributes)
                counters["accepted"] += 1
                counters["public_catalog"] += item["public_catalog"]
                counters["attributes"] += len(attributes)
                counters["with_attributes"] += bool(attributes)
                counters["unknown_family"] += item["classification"]["family_id"] == "unknown"
                counters["conflicting_attributes"] += sum(a["knowledge_state"] == "conflicting" for a in attributes)
                counters["unmapped_attributes"] += sum(not a["mapped"] for a in attributes)
                name = "own_items.jsonl" if policy["sources"][source]["kind"] == "own" else "competitor_items.jsonl"
                outputs[name].write(dumps(item) + "\n")
                for attr in attributes:
                    outputs["attribute_observations.jsonl"].write(dumps(attr) + "\n")
                    if not attr["mapped"]:
                        unmapped[(source, "attribute", attr["raw_name"])] += 1
                family = item["classification"]
                if family["family_id"] == "unknown":
                    unmapped[(source, "category", item["category_original"])] += 1
                slot = families.setdefault(family["family_id"], {"label": family["family_label"], "proximity": family["proximity"], "counts": {}, "public_counts": {}})
                slot["counts"][source] = slot["counts"].get(source, 0) + 1
                slot["public_counts"][source] = slot["public_counts"].get(source, 0) + int(item["public_catalog"])
                for issue in item["quality_issues"]:
                    issues[issue] += 1
                    severity = "technical_conflict" if "conflict" in issue or "ambiguous" in issue else "warning"
                    issue_writer.writerow([source, spreadsheet_safe(item["record_key"]), record["position"], issue, severity, True])
                if policy["sources"][source]["kind"] == "competitor":
                    for scope in ["portwell_public", "choice_public", "group_public", "portwell_internal", "choice_internal", "group_internal"]:
                        coverage_writer.writerow([run_id, item["item_id"], source, spreadsheet_safe(key), family["family_id"], scope, "partial", "not_assessed", False, "technical_profiles_pending_validation"])
                if source == "eco":
                    raw = record["raw"]
                    legacy_code = text_value(raw, "Cod_Material_Choice_Mais_Proximo")
                    legacy_writer.writerow([spreadsheet_safe(key), spreadsheet_safe(text_value(raw, "Resultado_Cruzamento")), spreadsheet_safe(legacy_code), bool(legacy_code and legacy_code in choice_skus), "not_assessed"])
            if counters["received"] != counters["accepted"] + counters["quarantine"]:
                raise ValueError("Reconciliação falhou")
            counts[source], quality[source] = dict(counters), dict(issues)
            print(f"{source}: {counters['received']} recebidos, {counters['accepted']} aceitos, {counters['quarantine']} quarentena", flush=True)
    finally:
        for stream in outputs.values():
            stream.close()
    for entry, info in zip(entries, inputs):
        if sha256_file(entry["path"]) != info["sha256"]:
            raise ValueError("Entrada mudou durante execução; publicação bloqueada")
    with (work / "taxonomy_unmapped.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["source_id", "kind", "label", "count"])
        writer.writerows((s, k, spreadsheet_safe(label), n) for (s, k, label), n in sorted(unmapped.items()))
    summary = {"schema_version": 1, "run_id": run_id, "mode": "refined_staging", "status": "base_validated",
               "created_at": datetime.now(timezone.utc).isoformat(), "source_observed_at": None,
               "rule_version": policy["rule_version"], "inputs": inputs, "counts": counts, "quality": quality,
               "families": families, "attribute_analysis": attribute_analysis.summary(), "catalog_completeness": "partial", "matching_status": "not_assessed",
               "strict_coverage_percent": None, "ai_calls": 0, "ai_cost": "0",
               "elapsed_seconds": round(perf_counter() - started, 3),
               "phases": {"ingestion": "complete", "normalization": "initial_rules_review_required", "matching": "pending", "opportunities": "pending", "incremental_review": "pending"}}
    atomic_json(work / "summary.json", summary)
    (work / "summary.md").write_text("# Nova base analisada\n\n" + "\n".join(f"- {s}: {n['accepted']} aceitos; {n.get('quarantine', 0)} em quarentena." for s, n in counts.items()) + "\n\nCatálogo próprio parcial. Famílias propostas; cobertura técnica não avaliada. Banco legado preservado. IA: zero chamadas.\n", encoding="utf-8")
    files = {p.name: {"sha256": sha256_file(p), "size_bytes": p.stat().st_size} for p in work.iterdir() if p.is_file() and p.name != "run_manifest.json"}
    manifest = {**summary, "output_files": files}
    atomic_json(work / "run_manifest.json", manifest)
    work.rename(target)
    verify_run(target)
    atomic_json(output_root / "latest.json", {"run_id": run_id})
    return {**manifest, "idempotent": False}


def verify_run(directory):
    directory = Path(directory)
    manifest = json.loads((directory / "run_manifest.json").read_text(encoding="utf-8"))
    if manifest.get("status") != "base_validated" or manifest.get("schema_version") != 1:
        raise ValueError("Execução não validada")
    for name, expected in manifest["output_files"].items():
        path = (directory / name).resolve()
        if path.parent != directory.resolve() or sha256_file(path) != expected["sha256"]:
            raise ValueError("Saída ausente ou alterada: " + name)
    return manifest


def publish_analysis(run_id, output_root=STAGING):
    if not re.fullmatch(r"[0-9a-f]{24}", run_id):
        raise ValueError("ID de execução inválido")
    output_root = Path(output_root)
    manifest = verify_run(output_root / run_id)
    if any(c.get("quarantine", 0) for c in manifest["counts"].values()):
        raise ValueError("Quarentena exige revisão antes da publicação")
    atomic_json(output_root / "published.json", {"run_id": run_id, "published_at": datetime.now(timezone.utc).isoformat(), "summary_sha256": manifest["output_files"]["summary.json"]["sha256"]})
    return manifest


def load_analysis(*, published=True, output_root=STAGING):
    try:
        root = Path(output_root)
        pointer = json.loads((root / ("published.json" if published else "latest.json")).read_text(encoding="utf-8"))
        run_id = pointer["run_id"]
        if not re.fullmatch(r"[0-9a-f]{24}", run_id):
            return None
        path = root / run_id / "summary.json"
        if path.stat().st_size > 2 * 1024 * 1024 or (published and sha256_file(path) != pointer.get("summary_sha256")):
            return None
        summary = json.loads(path.read_text(encoding="utf-8"))
        if summary.get("schema_version") != 1 or summary.get("status") != "base_validated" or summary.get("run_id") != run_id:
            return None
        return {**summary, "published": published}
    except (OSError, ValueError, KeyError, TypeError):
        return None
