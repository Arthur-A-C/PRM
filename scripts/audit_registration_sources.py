"""Read-only structural audit of supplied WooCommerce and CIGAM files."""

import csv
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path("C:/Users/PW/Downloads")
EXPORTS = (
    "wc-product-export-8-10-2026-1791468032367(in).csv",
    "wc-product-export-8-10-2026-1791468059447(in).csv",
    "wc-product-export-8-10-2026-1791468702941.csv",
)


def audit_csv(path):
    try:
        path.read_text(encoding="utf-8-sig")
        encoding = "utf-8-sig"
    except UnicodeDecodeError:
        encoding = "cp1252"
    with path.open("r", encoding=encoding, newline="") as source:
        reader = csv.DictReader(source)
        headers = reader.fieldnames or []
        count = 0
        sku_count = 0
        names = []
        types = Counter()
        sku_values = set()
        sku_header = next((name for name in headers if name.casefold() == "sku"), None)
        for row in reader:
            count += 1
            sku = row.get(sku_header, "") if sku_header else ""
            if sku:
                sku_count += 1
                sku_values.add(sku)
            types[row.get("Tipo", "")] += 1
            if len(names) < 3:
                names.append({"sku": sku, "name": row.get("Nome", "")[:100]})
    return {
        "filename": path.name, "bytes": path.stat().st_size, "encoding": encoding,
        "headers": headers, "rows": count, "sku_rows": sku_count,
        "distinct_skus": len(sku_values), "types": types,
        "sample": names,
    }


def audit_cigam(path):
    book = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    for sheet in book:
        rows = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 5), values_only=True))
        sheets.append({
            "name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
            "first_rows": [[str(value)[:80] if value is not None else None for value in row] for row in rows],
        })
    book.close()
    return {"filename": path.name, "sheets": sheets}


def main():
    print(json.dumps({
        "exports": [audit_csv(ROOT / name) for name in EXPORTS],
        "cigam": audit_cigam(ROOT / "cadastro-cigam.xlsx"),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
