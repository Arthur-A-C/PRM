import csv
import io
import json
import zipfile
from pathlib import Path

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image

from app.schemas.registration import RegistrationProduct
from app.services.registration_cigam import export_cigam_xlsx
from app.services.registration_core import category_levels, normalize_attributes, split_terms
from app.services.registration_description import build_description
from app.services.registration_exports import load_company_export
from app.services.registration_files import classify_files, find_product_folder
from app.services.registration_package import build_package
from app.services.registration_sheets import preview_sheet
from app.services.registration_woo_csv import export_woo_csv, parse_woo_export
from app.services.registration_woo_api import WooClient, WooCredentials


def product(**overrides):
    values = {
        "company": "choice", "cigam_code": "00123", "sku": "00123",
        "technical_name": "Nome técnico", "commercial_title": "Título comercial",
        "category_paths": category_levels("Industrial > Painéis"),
        "attributes": [{"name": "Tipo", "values": ["Painel"]}],
    }
    values.update(overrides)
    return RegistrationProduct(**values)


def test_product_codes_and_attribute_rules():
    assert product().sku == "00123"
    with pytest.raises(ValueError):
        product(sku=123)
    assert split_terms("0,4m, azul", multivalued=True) == ["0,4m", "azul"]
    attributes, warnings = normalize_attributes({
        "FrEquência": "50 Hz", "Tipo": "Painel", "tipo": "Painel",
        "Tensão": "-", "Observações": "ignorar",
    })
    assert [item.name for item in attributes] == ["Tipo", "Frequência"]
    assert warnings
    many, warnings = normalize_attributes({"Tipo": "A", **{f"Campo {i}": "x" for i in range(20)}})
    assert len(many) == 15
    assert "Campo" in " ".join(warnings)


def test_sheet_conflict_priority_and_duplicate_tabs():
    book = Workbook()
    first = book.active
    first.title = "Aba 2"
    for row in (
        ["Voltar ao Índice"], ["Industrial > Painéis | Responsável: Pessoa"],
        ["Nome do Produto", "Título", "Fornecedor", "Código CIGAM", "Status", "Tipo", "Observações"],
        ["Técnico", "Título antigo", "Fabricante", "00123", "Pendente", "Painel", "ignorar"],
    ):
        first.append(row)
    second = book.copy_worksheet(first)
    second.title = "Aba 3"
    third = book.create_sheet("Aba 1")
    for row in (
        ["Voltar ao Índice"], ["Industrial > Computadores | Responsável: Pessoa"],
        ["Nome do Produto", "Título", "Fornecedor", "Código CIGAM", "Status", "Tipo", "Materiais Disponíveis (link)"],
        ["Técnico", "Título aprovado", "Fabricante", "00123", "Concluído", "PC", "url"],
    ):
        third.append(row)
    output = io.BytesIO()
    book.save(output)
    result = preview_sheet(output.getvalue(), "cadastro.xlsx", "choice", {"00123"})
    assert result["distinct_products"] == 1
    assert result["duplicate_tabs"] == ["Aba 3"]
    assert result["products"][0]["technical_name"] == "Técnico"
    assert result["products"][0]["commercial_title"] == "Título aprovado"
    assert result["products"][0]["sku"] == "00123"
    assert result["existing_count"] == 1
    assert any("títulos divergentes" in item for item in result["conflicts"])


def test_woo_csv_preserves_unapproved_fields_and_column_order():
    headers = [
        "ID", "Tipo", "SKU", "Nome", "Publicado", "Descrição curta",
        "Descrição", "Categorias", "Imagens", "Preço",
        "Nome do atributo 1", "Valores do atributo 1",
        "Visibilidade do atributo 1", "Atributo global 1",
    ]
    old = ["42", "simple", "00123", "Antigo", "1", "Breve", "<p>Antigo</p>",
           "Industrial", "https://image", "999", "Tipo", "PC", "1", "1"]
    source = io.StringIO(newline="")
    writer = csv.writer(source)
    writer.writerow(headers)
    writer.writerow(old)
    template = parse_woo_export(source.getvalue().encode("utf-8"))
    result, warnings = export_woo_csv(
        product(attributes=[]), template, approved_fields={"title", "categories"}
    )
    rows = list(csv.DictReader(io.StringIO(result.decode("utf-8-sig"))))
    assert list(rows[0]) == headers
    assert rows[0]["ID"] == "42"
    assert rows[0]["Preço"] == "999"
    assert rows[0]["Imagens"] == "https://image"
    assert rows[0]["Publicado"] == "1"
    assert rows[0]["Nome"] == "Título comercial"
    assert rows[0]["Nome do atributo 1"] == "Tipo"
    assert not warnings


def test_woo_rejects_encapsulated_header():
    with pytest.raises(ValueError, match="WooCommerce"):
        parse_woo_export(b'"ID,Tipo,SKU,Nome";;;;;\n')


def test_woo_mixed_export_batches_and_ambiguous_global(monkeypatch, tmp_path):
    headers = ["ID", "Tipo", "SKU", "Nome", "Categorias", "Descrição curta", "Descrição",
               "Nome do atributo 1", "Valores do atributo 1", "Visibilidade do atributo 1",
               "Atributo global 1", "Metadado: lote"]
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerow(["1", "simple", "A", "A", "X", "", "", "Tipo", "X", "1", "1"])
    writer.writerow(["2", "simple", "B", "B", "X", "", "", "Tipo", "Y", "1", "0", "novo"])
    archive = tmp_path / "exports.zip"
    with zipfile.ZipFile(archive, "w") as package:
        package.writestr("export_woocommerce_choice.csv", output.getvalue())
    monkeypatch.setenv("PRM_WOO_EXPORTS_ZIP", str(archive))
    template = parse_woo_export(load_company_export("choice"))
    assert template.padded_rows == 1
    assert template.ambiguous_globals == {"tipo"}
    assert template.rows_by_sku["A"]["Metadado: lote"] == ""
    with pytest.raises(ValueError, match="global e local"):
        export_woo_csv(product(sku="C", cigam_code="C"), template)


def test_description_signature_and_escaping():
    html, warnings = build_description(product(benefits=["Sem <risco>"]))
    assert "<h1>Título comercial</h1>" in html
    assert "Sem &lt;risco&gt;" in html
    assert "<strong>Choice automação" in html
    assert not warnings
    other, warnings = build_description(product(company="portwell"))
    assert "Choice automação" not in other
    assert warnings


def test_cigam_template_and_real_zip():
    template = Workbook()
    template.active.append(["Modelo", "Codigo Cigam", "NCM"])
    output = io.BytesIO()
    template.save(output)
    content, warnings = export_cigam_xlsx(product(model="ABC"), output.getvalue())
    sheet = load_workbook(io.BytesIO(content), read_only=True).active
    assert list(sheet.values)[1] == ("ABC", "00123", None)
    assert warnings == ["Campo CIGAM pendente: NCM"]
    image = io.BytesIO()
    Image.new("RGB", (10, 20), "red").save(image, format="PNG")
    package, manifest = build_package(
        product(), [("original.png", image.getvalue(), 90)],
        [("source_datasheet", "original.pdf", b"%PDF-test")],
        size=1200, quality=90,
    )
    # The supplier datasheet is accepted as input but never exported or listed.
    with zipfile.ZipFile(io.BytesIO(package)) as archive:
        assert set(archive.namelist()) == {
            "00123/1.jpg", "00123/manifest.json",
        }
        processed = Image.open(io.BytesIO(archive.read("00123/1.jpg")))
        assert processed.size == (20, 10)
        assert json.loads(archive.read("00123/manifest.json"))["sku"] == "00123"
    assert [f["role"] for f in manifest["files"]] == ["main_image"]


def test_woo_live_preview_is_read_only_and_preserves_existing(monkeypatch):
    calls = []
    def answer(request):
        calls.append((request.method, request.url.path))
        if request.url.path.endswith("/products"):
            return httpx.Response(200, json=[{
                "id": 12, "sku": "00123", "name": "Anterior", "status": "publish",
                "categories": [{"id": 4}], "attributes": [
                    {"id": 7, "name": "Material", "options": ["Aço"], "visible": True}
                ],
            }])
        if request.url.path.endswith("/products/categories"):
            return httpx.Response(200, json=[
                {"id": 4, "name": "Industrial", "parent": 0},
                {"id": 5, "name": "Painéis", "parent": 4},
            ])
        if request.url.path.endswith("/products/attributes"):
            return httpx.Response(200, json=[{"id": 8, "name": "Tipo"}])
        return httpx.Response(404)
    client = WooClient(
        WooCredentials("https://example.test", "ck_test", "cs_test"),
        transport=httpx.MockTransport(answer),
    )
    try:
        result = client.preview(product())
    finally:
        client.close()
    assert result["action"] == "update"
    assert result["product_id"] == 12
    assert result["payload"]["categories"] == [{"id": 4}, {"id": 5}]
    assert len(result["payload"]["attributes"]) == 2
    assert result["payload"]["attributes"][1]["id"] == 8
    assert "description" not in result["payload"]
    assert "short_description" not in result["payload"]
    assert all(method == "GET" for method, _ in calls)


def test_file_discovery_uses_exact_sku_and_classifies(tmp_path):
    folder = tmp_path / "00123"
    folder.mkdir()
    (folder / "00123_variante.jpg").write_bytes(b"img")
    (folder / "00123_foto2.jpg").write_bytes(b"img")
    (folder / "00123_Datasheet.pdf").write_bytes(b"pdf")
    found = find_product_folder(tmp_path, "00123")
    assert found == folder
    classified = classify_files(found)
    assert classified["main_image"].endswith("00123_variante.jpg")
    assert len(classified["gallery"]) == 1
    assert classified["documents"][0]["role"] == "source_datasheet"
    with pytest.raises(ValueError):
        find_product_folder(tmp_path, "../00123")
