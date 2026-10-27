import uuid
from types import SimpleNamespace

import pytest
from fastapi import FastAPI, Depends
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, select, func
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.compiler import compiles

from app.config import Settings, get_settings
from app.db import get_db_session
from app.dependencies import get_current_auth, require_module_access
from app.models import (Organization, Company, User, Offer, CatalogImport, Brand, BrandAlias,
                        Category, ImportRow, OfferEnrichment, OfferObservation, OfferPrice,
                        AttributeDefinition, OfferAttribute)
from app.repositories.catalog import CatalogFilters, list_offers, comparison_preview
from app.routers import catalog_imports
from app.services.catalog_import import publish_catalog
from app.services.catalog_origin import resolve_origin
from app.services.catalog_table import read_table, mapped_preview, suggest_mapping
from app.services.security import hash_secret


@compiles(JSONB, "sqlite")
def sqlite_json(type_, compiler, **kw):
    return "JSON"


@pytest.fixture
def catalog_db():
    engine = create_engine("sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False})
    for model in (Organization, User, Company, Brand, BrandAlias, Category, CatalogImport, Offer,
                  OfferEnrichment, ImportRow, OfferObservation, OfferPrice, AttributeDefinition, OfferAttribute):
        model.__table__.create(engine)
    with Session(engine, expire_on_commit=False) as db:
        org = Organization(name="Organização de teste"); other = Organization(name="Outra organização")
        db.add_all([org, other]); db.flush()
        user = User(organization_id=org.id, name="Gestor", email="gestor@example.test", password_hash="test", role="gestor")
        db.add(user)
        companies = {}
        for code, kind, organization in [("portwell", "nossa", org), ("choice", "nossa", org),
                                          ("sag", "nossa", org), ("rival", "concorrente", org),
                                          ("rival2", "concorrente", org), ("outside", "nossa", other)]:
            c = Company(organization_id=organization.id, code=code, name=code, company_type=kind, active=True)
            db.add(c); companies[code] = c
        db.commit()
        yield SimpleNamespace(db=db, org=org, other=other, companies=companies, user=user)
    engine.dispose()


def preview(tmp_path, company, name="Produto", sku="001", suffix=""):
    path = tmp_path / (company.code + suffix + ".csv")
    path.write_text(f"SKU;Nome;Preco\n{sku};{name};1.250,00\n", encoding="utf-8-sig")
    table = read_table(path)
    return mapped_preview(path, company.code, table, suggest_mapping(table["headers"]))


def publish(state, tmp_path, code, name="Produto"):
    company = state.companies[code]
    result = publish_catalog(state.db, preview(tmp_path, company, name), company)
    state.db.commit()
    return result


def test_equal_skus_never_collide_and_updates_are_source_scoped(catalog_db, tmp_path):
    s = catalog_db
    for code in ("portwell", "choice", "rival", "rival2", "outside"):
        assert publish(s, tmp_path, code).inserted == 1
    assert s.db.scalar(select(func.count()).select_from(Offer)) == 5
    assert publish(s, tmp_path, "portwell", "Portwell atualizado").updated == 1
    assert publish(s, tmp_path, "rival", "Concorrente atualizado").updated == 1
    names = dict(s.db.execute(select(Company.code, Offer.name_raw).join(Offer)).all())
    assert names == {"portwell": "Portwell atualizado", "rival": "Concorrente atualizado",
                     "choice": "Produto", "rival2": "Produto", "outside": "Produto"}


def test_filters_and_manual_comparator_keep_group_and_tenant(catalog_db, tmp_path):
    s = catalog_db
    for code in ("portwell", "choice", "rival", "outside"):
        publish(s, tmp_path, code)
    own = list_offers(s.db, s.org.id, CatalogFilters(source_group="internal"))
    rivals = list_offers(s.db, s.org.id, CatalogFilters(source_group="competitor"))
    assert {i["company_code"] for i in own["items"]} == {"portwell", "choice"}
    assert {i["company_code"] for i in rivals["items"]} == {"rival"}
    comparison = comparison_preview(s.db, s.org.id, [own["items"][0]["id"], rivals["items"][0]["id"]])
    assert [i["company"]["source_group"] for i in comparison["items"]] == ["internal", "competitor"]
    # Intentional internal comparisons remain supported, without relabeling a company.
    comparison = comparison_preview(s.db, s.org.id, [i["id"] for i in own["items"]])
    assert all(i["company"]["source_group"] == "internal" for i in comparison["items"])
    outside = s.db.scalar(select(Offer.id).where(Offer.organization_id == s.other.id))
    with pytest.raises(LookupError):
        comparison_preview(s.db, s.org.id, [own["items"][0]["id"], outside])


def test_origin_rejects_wrong_group_or_organization(catalog_db):
    s = catalog_db
    for code, group in [("portwell", "competitor"), ("rival", "internal"), ("outside", "internal")]:
        with pytest.raises(LookupError):
            resolve_origin(s.db, s.org.id, s.companies[code].id, group)


@pytest.mark.parametrize("delimiter", [",", ";"])
def test_csv_mapping_preserves_sku_and_price(tmp_path, delimiter):
    path = tmp_path / "sample.csv"
    import csv
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file, delimiter=delimiter)
        writer.writerows([["SKU", "Nome", "Preco"], ["0001", "Produto", "1.250,00"], ["0001", "Duplicado", "12.50"]])
    table = read_table(path)
    result = mapped_preview(path, "any-company", table, suggest_mapping(table["headers"]))
    assert result.rows[0].external_code == "0001"
    assert str(result.rows[0].prices[0].amount) == "1250.00"
    assert result.rows[1].errors == ["SKU duplicado neste arquivo"]


def test_excel_requires_sheet_and_never_executes_formulas(tmp_path):
    path = tmp_path / "test.xlsx"
    workbook = Workbook()
    workbook.active.append(["SKU", "Nome"]); workbook.active.append(["001", "Produto"])
    workbook.create_sheet("Segunda").append(["SKU", "Nome"])
    workbook["Segunda"].append(["002", '=HYPERLINK("https://example.test", "unsafe")'])
    workbook.save(path)
    assert read_table(path)["headers"] == []
    first = read_table(path, sheet="Sheet")
    assert mapped_preview(path, "sag", first, suggest_mapping(first["headers"])).valid_rows == 1
    second = read_table(path, sheet="Segunda")
    assert mapped_preview(path, "sag", second, suggest_mapping(second["headers"])).rejected_rows == 1


@pytest.fixture
def client(catalog_db):
    s = catalog_db
    app = FastAPI(); app.include_router(catalog_imports.router, dependencies=[Depends(require_module_access)])
    auth = SimpleNamespace(user=s.user, login_session=SimpleNamespace(csrf_hash=hash_secret("csrf")))
    app.dependency_overrides[get_current_auth] = lambda: auth
    app.dependency_overrides[get_db_session] = lambda: s.db
    app.dependency_overrides[get_settings] = lambda: Settings(_env_file=None)
    return TestClient(app), auth


def post(client, company, action, **extra):
    import json
    data = {"company_id": str(company.id), "group": "internal", "action": action,
            "mapping": json.dumps({"external_code": "SKU", "name": "Nome"}), **extra}
    return client.post("/api/v1/catalog-imports/prepare", data=data,
                       files={"file": ("sample.csv", b"SKU;Nome\n001;Produto\n", "text/csv")},
                       headers={"X-CSRF-Token": "csrf"})


def test_api_preview_no_write_confirm_history_and_reimport(catalog_db, client):
    s = catalog_db; web, auth = client; c = s.companies["sag"]
    response = post(web, c, "preview"); assert response.status_code == 200, response.text
    assert s.db.scalar(select(func.count()).select_from(Offer)) == 0
    assert s.db.scalar(select(func.count()).select_from(CatalogImport)) == 0
    result = post(web, c, "import", expected_preview=response.json()["preview_token"])
    assert result.status_code == 200, result.text
    assert result.json()["result"]["inserted"] == 1
    h = web.get("/api/v1/catalog-imports/history").json()["items"][0]
    assert (h["source_group"], h["origin"], h["filename"], h["user"]) == ("internal", "sag", "sample.csv", "Gestor")
    token = post(web, c, "preview").json()["preview_token"]
    assert post(web, c, "import", expected_preview=token).json()["result"]["idempotent"]


def test_api_rejects_stale_preview_and_unauthorized(catalog_db, client, tmp_path):
    s = catalog_db; web, auth = client; c = s.companies["portwell"]
    token = post(web, c, "preview").json()["preview_token"]
    publish(s, tmp_path, "portwell", "Alteração concorrente")
    assert post(web, c, "import", expected_preview=token).status_code == 409
    assert post(web, s.companies["outside"], "preview").status_code == 404
    assert post(web, c, "preview", group="competitor").status_code == 404
    auth.user.role = "vendedor"
    assert post(web, c, "preview").status_code == 403


def test_import_modes_do_not_change_other_sources(catalog_db, tmp_path):
    s = catalog_db; c = s.companies["portwell"]
    publish(s, tmp_path, "portwell")
    r = publish_catalog(s.db, preview(tmp_path, c, "Não substituir"), c, mode="add_only")
    s.db.commit(); assert r.updated == 0 and r.unchanged == 1
    r = publish_catalog(s.db, preview(tmp_path, c, "Não criar", "002"), c, mode="update_only")
    s.db.commit(); assert r.inserted == 0
    assert s.db.scalar(select(Offer.name_raw)) == "Produto"


def test_automatic_comparator_only_lists_competitors(catalog_db):
    from app.models.matching import CatalogMatchRun
    from app.repositories.matching import catalog_match_report
    s = catalog_db
    CatalogMatchRun.__table__.create(s.db.get_bind())
    report = catalog_match_report(s.db, s.org.id)
    assert {c["code"] for c in report["companies"]} == {"rival", "rival2"}


def test_competitor_creation_and_csrf(catalog_db, client):
    s = catalog_db; web, auth = client
    payload = {"name": "Empresa XYZ", "domain": "https://example.test"}
    assert web.post("/api/v1/catalog-imports/competitors", json=payload).status_code == 403
    response = web.post("/api/v1/catalog-imports/competitors", json=payload, headers={"X-CSRF-Token": "csrf"})
    assert response.status_code == 201
    c = s.db.get(Company, uuid.UUID(response.json()["id"]))
    assert c.company_type == "concorrente" and c.organization_id == s.org.id
    r = post(web, c, "preview", group="competitor"); assert r.status_code == 200
    r = post(web, c, "import", group="competitor", expected_preview=r.json()["preview_token"])
    assert r.status_code == 200 and r.json()["result"]["inserted"] == 1


def test_failed_import_rolls_back(catalog_db, client, monkeypatch):
    s = catalog_db; web, auth = client; c = s.companies["portwell"]
    token = post(web, c, "preview").json()["preview_token"]
    original = catalog_imports.publish_catalog
    def fail(*args, **kwargs):
        original(*args, **kwargs)
        args[0].flush()
        raise ValueError("Falha simulada")
    monkeypatch.setattr(catalog_imports, "publish_catalog", fail)
    assert post(web, c, "import", expected_preview=token).status_code == 422
    assert s.db.scalar(select(func.count()).select_from(Offer)) == 0
    assert s.db.scalar(select(func.count()).select_from(CatalogImport)) == 0


def test_bad_excel_is_reported_without_import(catalog_db, client):
    web, auth = client; c = catalog_db.companies["portwell"]
    response = web.post("/api/v1/catalog-imports/prepare", data={"company_id": str(c.id), "group": "internal"},
                        files={"file": ("bad.xlsx", b"not excel")}, headers={"X-CSRF-Token": "csrf"})
    assert response.status_code == 422


@pytest.mark.parametrize('extension', ['csv', 'xlsx'])
def test_csrf_protects_inspect_preview_and_confirmation(catalog_db, client, extension):
    from io import BytesIO
    import json
    s = catalog_db; web, auth = client
    content = b'SKU;Nome\n001;Produto\n'
    if extension == 'xlsx':
        workbook = Workbook(); workbook.active.append(['SKU', 'Nome'])
        workbook.active.append(['001', 'Produto'])
        buffer = BytesIO(); workbook.save(buffer); workbook.close(); content = buffer.getvalue()
    data = {'company_id': str(s.companies['portwell'].id), 'group': 'internal',
            'mapping': json.dumps({'external_code': 'SKU', 'name': 'Nome'})}
    def request(action, token, **extra):
        return web.post('/api/v1/catalog-imports/prepare', data={**data, 'action': action, **extra},
                        files={'file': ('products.' + extension, content)},
                        headers={'X-CSRF-Token': token} if token is not None else {})
    for action in ('inspect', 'preview', 'import'):
        for token in (None, 'wrong-token'):
            response = request(action, token)
            assert response.status_code == 403 and response.json()['detail'] == 'Token CSRF inválido.'
    inspected = request('inspect', 'csrf')
    assert inspected.status_code == 200 and inspected.json()['headers'] == ['SKU', 'Nome']
    preview = request('preview', 'csrf')
    assert preview.status_code == 200
    assert s.db.scalar(select(func.count()).select_from(Offer)) == 0
    assert s.db.scalar(select(func.count()).select_from(CatalogImport)) == 0
    expected = preview.json()['preview_token']
    assert request('import', None, expected_preview=expected).status_code == 403
    assert s.db.scalar(select(func.count()).select_from(Offer)) == 0
    published = request('import', 'csrf', expected_preview=expected)
    assert published.status_code == 200 and published.json()['result']['inserted'] == 1


def test_published_snapshot_filters_preserve_tenant(catalog_db):
    from app.models.refined_catalog import RefinedCatalogRun, RefinedCatalogItem
    from app.services.refined_catalog_database import list_refined_items, current_items
    s = catalog_db
    for model in (RefinedCatalogRun, RefinedCatalogItem):
        model.__table__.create(s.db.get_bind())
    for org, run in ((s.org, 'a' * 24), (s.other, 'b' * 24)):
        s.db.add(RefinedCatalogRun(id=run, organization_id=org.id, status='published', manifest={'run_id': run, 'families': {}}))
        for source in ('portwell', 'choice', 'eco', 'dimensional'):
            identifier = uuid.uuid4()
            s.db.add(RefinedCatalogItem(run_id=run, item_id=identifier, source_id=source, record_key='001',
                                       title_search='produto', family_id='unknown', public_catalog=True,
                                       payload={'item_id': str(identifier), 'source_id': source, 'title_search': 'produto'}))
    s.db.commit()
    own = list_refined_items(s.db, s.org.id, source_group='internal')
    rivals = list_refined_items(s.db, s.org.id, source_group='competitor')
    assert own['total'] == 2 and rivals['total'] == 2
    assert {i['source_id'] for i in own['items']} == {'portwell', 'choice'}
    assert {i['source_id'] for i in rivals['items']} == {'eco', 'dimensional'}
    assert all(i['source_group'] == 'internal' for i in own['items'])
    outside = list_refined_items(s.db, s.other.id)['items'][0]['item_id']
    with pytest.raises(LookupError):
        current_items(s.db, s.org.id, [uuid.UUID(outside)])


def test_blank_lines_keep_original_numbers_and_nonfinite_prices_are_invalid(tmp_path):
    path = tmp_path / 'test.csv'
    path.write_text('SKU;Nome;Preco\n\n001;Produto;NaN\n', encoding='utf-8')
    table = read_table(path)
    result = mapped_preview(path, 'choice', table, suggest_mapping(table['headers']))
    assert result.rows[0].row_number == 3
    assert result.rows[0].errors == ['Preço venda inválido']
