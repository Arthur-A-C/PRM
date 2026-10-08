"""Regressões ancoradas nos registros reais; os originais são somente leitura."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services.catalog_semantics import analyze_technical, assess_pair


IMPORTS = Path(__file__).resolve().parents[1] / "data" / "imports"


def row_at(filename: str, line: int):
    path = IMPORTS / filename
    if not path.is_file():
        pytest.skip(f"fixture ausente: {path.name}")
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook.active
        sheet.reset_dimensions()
        return next(sheet.iter_rows(min_row=line, max_row=line, values_only=True))
    finally:
        workbook.close()


def test_real_portwell_892_does_not_turn_ssd_into_screen():
    row = row_at("produtos-portwell.xlsx", 13)
    assert row[4] == "316001000892"
    result = analyze_technical(name=row[7], application=row[8])
    assert result.form_factor == "box_pc"
    assert "Intel Core i7-13700TE" in result.values("cpu.model")
    assert any(device["quantity"] == 2 and device["capacity"] == "1" and device["form_factor_in"] == "2.5" for device in result.values("storage.devices"))
    assert not result.values("display.diagonal_in")
    assert result.values("ports.usb") == [{"quantity": 6, "version": "3.2"}, {"quantity": 2, "version": "2.0"}]


def test_real_choice_classification_conflicts_are_review_only():
    ambiguous = row_at("produtos-choice.xlsx", 2868)
    assert ambiguous[4] == "111001000409"
    result = analyze_technical(name=ambiguous[7], application=ambiguous[8], category=ambiguous[3])
    assert "revisar_subtipo_rele_contator" in result.issues
    assert not result.values("safety.certification")
    contactor = row_at("produtos-choice.xlsx", 10045)
    assert contactor[4] == "190003000491"
    choice = analyze_technical(name=contactor[7], application=contactor[8])
    assert choice.family == "contactor"
    assert {"value": 25, "utilization": "AC-3"} in choice.values("current.rating_a")
    eco = row_at("produtos_ecoautomacao.xlsx", 392)
    assert eco[0] == "P002212"
    competitor = analyze_technical(name=eco[1], description=eco[7])
    assert assess_pair(choice, competitor)["status"] != "identidade_confirmada"


def test_real_power_supply_conflict_blocks_legacy_high_score():
    eco = row_at("produtos_ecoautomacao.xlsx", 3146)
    choice = row_at("produtos-choice.xlsx", 3807)
    assert (eco[0], choice[4]) == ("P000194", "112008000051")
    competitor = analyze_technical(name=eco[1], description=eco[7])
    own = analyze_technical(name=choice[7], application=choice[8])
    assert assess_pair(competitor, own)["status"] == "incompativel"
